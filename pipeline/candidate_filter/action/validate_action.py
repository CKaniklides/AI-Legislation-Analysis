"""
Quick validation of the action candidate score.

Two steps:
  1. sample  - draw ~50 norm pairs across all score levels and write them to an
               Excel sheet for manual labelling (the scores are hidden from you).
  2. review  - read the filled-in sheet back and report how well the scores
               agree with your labels.

This is a confidence check, not a full evaluation: the sample is small and
stratified by score, so treat the numbers as indicative.

Usage:
  python validate_actions.py sample --norms AI_ACT_norms.json --out action_validation.xlsx
  python validate_actions.py review --xlsx action_validation.xlsx
"""

from __future__ import annotations

import argparse
import itertools
import random
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from action_scoring import build_profiles_from_norms, score_action_pair
from action_verbs import inner_norm, load_norm_records

# How many pairs to draw per stratum (total 50).
STRATA = {
    "1.0 identical lemma": 5,
    "1.0 same axis, different verb": 7,
    "1.0 same axis, generic verb": 4,
    "0.5 related axes": 14,
    "-0.5 unrelated": 12,
    "0.0 missing": 8,
}
MAX_USES_PER_NORM = 2

LABELS = ["same", "related", "unrelated", "unclear"]
LABEL_SCORE = {"same": 1.0, "related": 0.5, "unrelated": -0.5}

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
FILL_INPUT = PatternFill("solid", start_color="FFFF00")
FILL_HEAD = PatternFill("solid", start_color="D9D9D9")
WRAP = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------------
# Sampling
# ---------------------------------------------------------------------

def stratum(result, lemma_a, lemma_b) -> str | None:
    s, rel = result.action_score, result.action_relation
    if s == 1.0:
        if lemma_a == lemma_b:
            return "1.0 identical lemma"
        return "1.0 same axis, generic verb" if result.generic else "1.0 same axis, different verb"
    if s == 0.5:
        return "0.5 related axes"
    if s == -0.5:
        return "-0.5 unrelated"
    if rel == "missing" and (lemma_a or lemma_b):
        return "0.0 missing"
    return None


def excerpt(record: dict, width: int = 160) -> str:
    """Short piece of the paragraph around the action span."""
    norm = inner_norm(record)
    text = (record.get("text") or {}).get("paragraph") or ""
    span = [s for s in norm.get("action_span") or [] if isinstance(s, str) and s.strip()]
    pos = text.find(span[0]) if span else -1
    if pos < 0:
        return text[: 2 * width] + ("…" if len(text) > 2 * width else "")
    start, end = max(0, pos - width), min(len(text), pos + len(span[0]) + width)
    return ("…" if start else "") + text[start:end] + ("…" if end < len(text) else "")


def sample(norms_path: Path, out_path: Path, seed: int) -> None:
    records = {str(r.get("norm_id")): r for r in load_norm_records(norms_path)}
    profiles = build_profiles_from_norms(norms_path)
    ids = list(profiles)

    buckets: dict[str, list] = defaultdict(list)
    for a, b in itertools.combinations(ids, 2):
        pa, pb = profiles[a], profiles[b]
        result = score_action_pair(pa, pb)
        key = stratum(result, pa.lemma, pb.lemma)
        if key:
            buckets[key].append((a, b, result))

    rng = random.Random(seed)
    uses: Counter[str] = Counter()
    rows = []
    for key, n in STRATA.items():
        pool = buckets.get(key, [])
        rng.shuffle(pool)
        taken = 0
        for a, b, result in pool:
            if taken == n:
                break
            if uses[a] >= MAX_USES_PER_NORM or uses[b] >= MAX_USES_PER_NORM:
                continue
            uses[a] += 1
            uses[b] += 1
            taken += 1
            rows.append((key, a, b, result))
    rng.shuffle(rows)   # so the order doesn't reveal the score

    write_workbook(rows, records, profiles, out_path)
    print(f"Wrote {len(rows)} pairs -> {out_path}")
    print("Per stratum:", dict(Counter(r[0] for r in rows)))


def write_workbook(rows, records, profiles, out_path: Path) -> None:
    wb = Workbook()

    # --- Instructions ---------------------------------------------------------
    ins = wb.active
    ins.title = "Instructions"
    lines = [
        ("How to fill in", BOLD),
        ("Go to the 'Label' sheet. For each row, read the two actions and choose a label in the "
         "yellow column 'your_label' (dropdown). The 'notes' column is optional.", FONT),
        ("", FONT),
        ("Question: are the two norms about the same kind of action, ignoring negation and "
         "whether it is an obligation, prohibition or permission?", BOLD),
        ("same       = the same act, or its direct opposite (bewaren / verwijderen counts as same)", FONT),
        ("related    = neighbouring acts (verstrekken / melden, vastleggen / bewaren)", FONT),
        ("unrelated  = different acts (verstrekken / beoordelen)", FONT),
        ("unclear    = cannot tell from the text (excluded from the scores)", FONT),
        ("", FONT),
        ("Example (format only, not part of the sample):", BOLD),
        ("lemma_a: 'in kennis stellen'   lemma_b: 'op de hoogte brengen'   your_label: same", FONT),
        ("", FONT),
        ("The scores are hidden on purpose; don't look at the 'key' sheet until you're done. "
         "When finished, save the file and run: python validate_actions.py review --xlsx <file>", FONT),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        cell = ins.cell(row=i, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(wrap_text=True)
    ins.column_dimensions["A"].width = 110

    # --- Label sheet ----------------------------------------------------------
    ws = wb.create_sheet("Label")
    headers = ["row", "lemma_a", "span_a", "text_a", "lemma_b", "span_b", "text_b",
               "your_label", "notes", "norm_id_a", "norm_id_b"]
    widths = [6, 22, 22, 60, 22, 22, 60, 14, 30, 24, 24]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font, cell.fill = BOLD, FILL_HEAD
        ws.column_dimensions[cell.column_letter].width = w

    for i, (key, a, b, result) in enumerate(rows, start=1):
        na, nb = inner_norm(records[a]), inner_norm(records[b])
        values = [
            i,
            profiles[a].lemma or "(empty)", ", ".join(na.get("action_span") or []), excerpt(records[a]),
            profiles[b].lemma or "(empty)", ", ".join(nb.get("action_span") or []), excerpt(records[b]),
            None, None, a, b,
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=i + 1, column=col, value=v)
            cell.font, cell.alignment = FONT, WRAP
        ws.cell(row=i + 1, column=8).fill = FILL_INPUT

    dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H2:H{len(rows) + 1}")
    ws.freeze_panes = "B2"

    # --- Hidden key -----------------------------------------------------------
    key_ws = wb.create_sheet("key")
    key_headers = ["row", "stratum", "action_score", "action_relation", "generic", "basis"]
    for col, h in enumerate(key_headers, start=1):
        key_ws.cell(row=1, column=col, value=h).font = BOLD
    for i, (key, a, b, result) in enumerate(rows, start=1):
        for col, v in enumerate([i, key, result.action_score, result.action_relation,
                                 result.generic, result.basis], start=1):
            key_ws.cell(row=i + 1, column=col, value=v).font = FONT
    key_ws.sheet_state = "hidden"

    wb.active = 1
    wb.save(out_path)


# ---------------------------------------------------------------------
# Review
# ---------------------------------------------------------------------

def review(xlsx_path: Path) -> None:
    wb = load_workbook(xlsx_path, data_only=True)
    labels = {}
    for row in wb["Label"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            labels[int(row[0])] = (str(row[7]).strip().lower() if row[7] else None,
                                   row[1], row[4])
    key = {}
    for row in wb["key"].iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            key[int(row[0])] = {"stratum": row[1], "score": float(row[2])}

    labelled = [(i, *labels[i], key[i]) for i in sorted(labels)
                if labels[i][0] in LABEL_SCORE and i in key]
    unclear = sum(1 for v in labels.values() if v[0] == "unclear")
    blank = sum(1 for v in labels.values() if v[0] is None)
    print(f"Labelled: {len(labelled)}   unclear: {unclear}   blank: {blank}\n")
    if not labelled:
        print("Nothing to review yet - fill in 'your_label' first.")
        return

    # 1. Exact agreement on scored pairs (missing pairs have no claim to check).
    scored = [x for x in labelled if x[4]["score"] != 0.0]
    exact = sum(LABEL_SCORE[x[1]] == x[4]["score"] for x in scored)

    # 2. Candidate decision: does score >= 0.5 match "same or related"?
    cand_tp = sum(x[4]["score"] >= 0.5 and x[1] != "unrelated" for x in scored)
    cand_fp = sum(x[4]["score"] >= 0.5 and x[1] == "unrelated" for x in scored)
    cand_fn = sum(x[4]["score"] < 0.5 and x[1] != "unrelated" for x in scored)
    cand_tn = sum(x[4]["score"] < 0.5 and x[1] == "unrelated" for x in scored)
    cand_acc = (cand_tp + cand_tn) / len(scored) if scored else 0.0
    precision = cand_tp / (cand_tp + cand_fp) if cand_tp + cand_fp else 0.0

    print("Per stratum (share of your labels that agree with the score):")
    by_stratum = defaultdict(list)
    for x in labelled:
        by_stratum[x[4]["stratum"]].append(x)
    for name in STRATA:
        xs = by_stratum.get(name, [])
        if not xs:
            continue
        dist = Counter(x[1] for x in xs)
        if name.startswith("0.0"):
            hit = sum(x[1] == "unrelated" for x in xs)
            note = "labelled unrelated (higher = fewer missed candidates)"
        else:
            target = float(name.split()[0])
            hit = sum(LABEL_SCORE[x[1]] == target for x in xs)
            note = "exact match"
        print(f"  {name:32} {hit}/{len(xs)} {note:<10}  {dict(dist)}")

    print()
    print(f"Exact agreement (scored pairs):        {exact}/{len(scored)} = {exact / len(scored):.0%}")
    print(f"Candidate decision accuracy (>=0.5):   {cand_acc:.0%}  "
          f"[TP {cand_tp}, FP {cand_fp}, FN {cand_fn}, TN {cand_tn}]")
    print(f"Candidate precision:                   {precision:.0%}")

    missed = [x for x in labelled if x[4]["score"] == 0.0 and x[1] != "unrelated"]
    if missed:
        print(f"Missing-score pairs you judged same/related: {len(missed)} "
              "(lexicon gaps worth fixing)")

    verdict = ("good" if cand_acc >= 0.75 else "fair" if cand_acc >= 0.60 else "weak")
    print(f"\nConfidence: {verdict.upper()}  (good >= 75%, fair >= 60% on the candidate decision)")

    disagreements = [x for x in scored
                     if (x[4]["score"] >= 0.5) != (x[1] != "unrelated")]
    if disagreements:
        print("\nDisagreements on the candidate decision:")
        for i, label, la, lb, k in disagreements:
            print(f"  row {i:2}: {la!r} vs {lb!r}  score {k['score']:+.1f}, you: {label}")


# ---------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(description="Quick validation of action scores.")
    sub = parser.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample", help="Write a labelling sheet.")
    s.add_argument("--norms", type=Path, default=Path("data/conflict/original_norms.json"))
    s.add_argument("--out", type=Path, default=Path("data/conflict/action/action_validation.xlsx"))
    s.add_argument("--seed", type=int, default=13)
    r = sub.add_parser("review", help="Score a filled-in labelling sheet.")
    r.add_argument("--xlsx", type=Path, default=Path("action_validation.xlsx"))
    args = parser.parse_args()

    if args.cmd == "sample":
        sample(args.norms, args.out, args.seed)
    else:
        review(args.xlsx)


if __name__ == "__main__":
    main()