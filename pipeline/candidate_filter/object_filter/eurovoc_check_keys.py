"""
Quick check of the object keys (Step 1). Builds a small workbook:

  Heads   a random sample of units: is head_word the noun that says what the
          object is? Fill in judgement (correct / wrong).
  Merges  keys shared by different head words. Most are singular/plural (fine);
          look for different words merged by the stemmer (handel / handeling).
  Splits  the most frequent head words, alphabetically, with their key.
          Variants of one word sit next to each other; 'check' marks a word
          that starts like the previous one but has a different key
          (advies / adviezen).
  Tails   every head -> compound tail link. Fill in judgement (correct / wrong).
  Summary share correct for Heads and Tails (formulas).

Usage
  python check_keys.py
  python check_keys.py --dir data/conflict/eurovoc --n-heads 50 --seed 0
Output: <dir>/key_check.xlsx
"""

from __future__ import annotations

import argparse
import json
import random
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

INPUT = PatternFill("solid", start_color="FFF2CC")
BOLD = Font(bold=True)


def sheet(wb, title, headers, rows, judge_col=None, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = BOLD
    for r in rows:
        ws.append(r)
    if judge_col:
        col = chr(ord("A") + headers.index(judge_col))
        dv = DataValidation(type="list", formula1='"correct,wrong"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{len(rows) + 1}")
        for i in range(2, len(rows) + 2):
            ws[f"{col}{i}"].fill = INPUT
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    return ws


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dir", type=Path, default=Path("data/conflict/eurovoc"))
    p.add_argument("--n-heads", type=int, default=50)
    p.add_argument("--top-words", type=int, default=100)
    p.add_argument("--min-norms", type=int, default=3,
                   help="Merges: only keys that occur in at least this many norms")
    p.add_argument("--seed", type=int, default=0)
    args = p.parse_args()

    with open(args.dir / "unit_keys.jsonl", encoding="utf-8") as f:
        units = [json.loads(l) for l in f if l.strip()]
    with open(args.dir / "norm_keys.jsonl", encoding="utf-8") as f:
        norms = [json.loads(l) for l in f if l.strip()]

    # norms per head key and per head word
    key_norms, word_norms = Counter(), Counter()
    for n in norms:
        for k in {u["head"] for u in n["units"] if u["head"]}:
            key_norms[k] += 1
        for w in {u["head_word"] for u in n["units"] if u["head_word"]}:
            word_norms[w] += 1
    words_of_key = defaultdict(set)
    for u in units:
        if u["head"]:
            words_of_key[u["head"]].add(u["head_word"])

    wb = Workbook()
    wb.remove(wb.active)

    # Heads: random sample
    rng = random.Random(args.seed)
    sample = rng.sample(units, min(args.n_heads, len(units)))
    sheet(wb, "Heads", ["unit", "head_word", "head_key", "judgement", "right_head"],
          [[u["unit"], u["head_word"], u["head"], None, None] for u in sample],
          judge_col="judgement", widths={"A": 70, "B": 22, "C": 18, "D": 12, "E": 22})

    # Merges: one key, several head words
    merges = sorted(((k, sorted(ws), key_norms[k]) for k, ws in words_of_key.items()
                     if len(ws) > 1 and key_norms[k] >= args.min_norms),
                    key=lambda x: -x[2])
    sheet(wb, "Merges", ["key", "norms", "head_words"],
          [[k, n, ", ".join(ws)] for k, ws, n in merges],
          widths={"A": 18, "B": 8, "C": 80})

    # Splits: frequent head words alphabetically, flag neighbours with other keys
    word_key = {u["head_word"]: u["head"] for u in units if u["head_word"]}
    top = sorted(w for w, _ in word_norms.most_common(args.top_words))
    rows, prev = [], None
    for w in top:
        flag = "check" if prev and w[:4] == prev[:4] and word_key[w] != word_key[prev] else None
        rows.append([w, word_key[w], word_norms[w], flag])
        prev = w
    sheet(wb, "Splits", ["head_word", "key", "norms", "flag"], rows,
          widths={"A": 26, "B": 22, "C": 8, "D": 8})

    # Tails: every distinct head word -> tail link
    tails = Counter((u["head_word"], u["tail"]) for u in units if u["tail"])
    sheet(wb, "Tails", ["head_word", "tail", "units", "judgement"],
          [[h, t, n, None] for (h, t), n in sorted(tails.items(), key=lambda x: x[0][1])],
          judge_col="judgement", widths={"A": 40, "B": 18, "C": 8, "D": 12})

    # Summary
    ws = wb.create_sheet("Summary", 0)
    ws.append(["check", "judged", "correct", "share correct", "target"])
    for c in ws[1]:
        c.font = BOLD
    for name, col, target in [("Heads", "D", "~90%"), ("Tails", "D", "~90%")]:
        r = ws.max_row + 1
        ws.append([name, f'=COUNTA({name}!{col}2:{col}1000)',
                   f'=COUNTIF({name}!{col}2:{col}1000,"correct")',
                   f'=IFERROR(C{r}/B{r},"")', target])
        ws[f"D{r}"].number_format = "0%"
    ws.append([])
    ws.append(["Merges and Splits have no judgement column: scan them and note any "
               "words that are wrongly merged or split."])
    ws.column_dimensions["A"].width = 14
    for col in "BCDE":
        ws.column_dimensions[col].width = 14

    out = args.dir / "key_check.xlsx"
    wb.save(out)
    print(f"Wrote {out}: {len(sample)} heads, {len(merges)} merges, "
          f"{len(rows)} frequent words ({sum(1 for r in rows if r[3])} flagged), "
          f"{len(tails)} tails")


if __name__ == "__main__":
    main()