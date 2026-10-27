#!/usr/bin/env python3
"""actor_pair_review.py - validation of the actor pair scores (actor filter, Step 4).

Two steps:
  1. sample  - draw a stratified sample of scored Stage A pairs and write a workbook
               for blind labelling: do the two norms address the same actor?
  2. review  - read the labels back and report agreement with the score levels, the
               overlap decision (score >= 0.5), missed overlaps, the direction flag
               and the kinds of error, with 95% intervals.

Population
  The Stage A pairs that actor_scores.py scored (actor_pair_scores.jsonl, score != 0).
  Pairs scoring 0 make no claim and are not sampled: their numbers are exact counts,
  split into pairs where a norm has no usable actor (by the norm's status) and
  'unknown' pairs (both norms have an actor, but no relation is known).

Strata (justified in the workbook's 'Design' sheet), default allocation 120
  same | both exact                  15    weak | siblings                    12
  same | shortened match             10    weak | broad covers mapped         18
  nested                             20    weak | both broad, same class      10
  different | same document          18    different | other documents      17
  The relation behind each pair comes from ActorScorer.norm_pair (actor_scores.py,
  with the same relations and inventory as the scoring run). If that is not
  available, it is derived from norm_actor_map.jsonl and marked 'approx'.
  A norm appears in at most --max-uses sampled pairs (default 2).
  Optional: --check-unknown N draws N unknown EU-NL pairs, reported separately.

Labels: same / nested / related / different / unclear, plus 'Narrower actor?' for
pairs labelled same or nested (checks the direction flag).

Usage
  python actor_pair_review.py sample --data data/conflict/actor \
      --stage-a data/conflict/stageA_pairs.json --norm-text data/conflict/norms_all.json
  python actor_pair_review.py review --xlsx data/conflict/actor/actor_pair_review.xlsx
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import math
import random
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

LABELS = ["same", "nested", "related", "different", "unclear"]
NARROWER = ["A narrower", "B narrower", "equal", "can't say"]
LEVEL_OF_SCORE = {1.0: "same", 0.7: "nested", 0.2: "weak", -0.25: "different"}
EXPECTED = {"same": "same", "nested": "nested", "weak": "related", "different": "different"}
OVERLAP = {"same", "nested"}
FLAG_ANSWER = {"a_narrower": "A narrower", "b_narrower": "B narrower"}

STRATA = [  # (name, level, default n, why)
    ("same | both exact", "same", 15,
     "Most same pairs: both norms map to the same ID by an exact match. Checks the mapping of the "
     "common actors."),
    ("same | shortened match", "same", 10,
     "Same ID reached by dropping words from at least one actor: the riskier route to 'same'."),
    ("nested", "nested", 20,
     "One actor is a kind or part of the other (subtype, part_of). Also used to check the "
     "direction flag."),
    ("weak | siblings", "weak", 12,
     "Two kinds of one actor (and curated can_be_same / linked relations)."),
    ("weak | broad covers mapped", "weak", 18,
     "A broad (fallback) actor whose class fits the other's mapped actor: the largest weak cell."),
    ("weak | both broad, same class", "weak", 10,
     "Two broad actors of the same class: the weakest evidence."),
    ("different | same document", "different", 18,
     "Both mapped, no relation, norms from one document: the likeliest place to miss a relation."),
    ("different | other documents", "different", 17,
     "Both mapped, no relation, IDs valid in a shared document but norms from different documents."),
]
STRATUM_LEVEL = {s: lvl for s, lvl, _, _ in STRATA}
CHECK_UNKNOWN = "check | unknown EU-NL"

Z = 1.96
FONT = "Arial"
INPUT_FILL = PatternFill("solid", start_color="FFFF00")
HEAD_FILL = PatternFill("solid", start_color="1F3864")
THIN = Side(style="thin", color="BFBFBF")


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------

def _get(obj, dotted):
    for k in dotted.split("."):
        if not isinstance(obj, dict):
            return None
        obj = obj.get(k)
    return obj


def read_jsonl(path):
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def read_records(path):
    raw = Path(path).read_text(encoding="utf-8").strip()
    return json.loads(raw) if raw.startswith("[") else \
        [json.loads(l) for l in raw.splitlines() if l.strip()]


def source_of(norm_id):
    return "NL" if str(norm_id).upper().startswith("BWBR") else "EU"


def usable_actors(rec):
    return [a for a in rec.get("actors") or [] if a.get("actor_id") or a.get("fallback_class")]


def actor_key(a):
    return a["actor_id"] if a.get("actor_id") else f"class:{a['fallback_class']}"


def load_scorer(args, data, norm_path):
    """ActorScorer built like the scoring run, or None (then details are approximated)."""
    path = Path(args.scorer) if args.scorer else Path(__file__).with_name("actor_scores.py")
    if not path.is_file():
        print(f"Note: {path} not found; relations are derived from the actor map ('approx').")
        return None
    sys.path.insert(0, str(path.parent))   # actor_scores imports actor_lexicon from its folder
    spec = importlib.util.spec_from_file_location("actor_scores", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    rel_path = data / args.relations
    if not rel_path.is_file() or not rel_path.read_text(encoding="utf-8").strip():
        print(f"Note: {rel_path} not found or empty; relations are derived from the actor map ('approx').")
        return None
    cfg_path = data / "actor_scores_config.json"
    cfg = json.loads(cfg_path.read_text(encoding="utf-8")) if cfg_path.is_file() else {}
    if args.curated is None and cfg.get("curated"):
        args.curated = cfg["curated"]
        print(f"Using the curated relations of the scoring run: {args.curated}")
    if getattr(mod, "in_scope", None) is None:
        print("Note: actor_lexicon could not be imported by actor_scores.py; 'different' versus "
              "'unknown' may then differ from the scoring run.")
    curated = (data / args.curated) if args.curated else None
    inv = data / args.inventory
    entries = read_jsonl(inv) if inv.is_file() else []
    return mod.ActorScorer(mod.load_norm_actors(norm_path), mod.load_relations(rel_path, curated),
                           entries)


# ---------------------------------------------------------------------------
# Strata
# ---------------------------------------------------------------------------

def approx_detail(level, ra, rb):
    """Relation detail from the actor map alone, for when ActorScorer is unavailable."""
    xa, xb = usable_actors(ra), usable_actors(rb)
    ka, kb = {actor_key(a): a for a in xa}, {actor_key(b): b for b in xb}
    if level == "same":
        shared = sorted(k for k in set(ka) & set(kb) if not k.startswith("class:"))
        k = shared[0] if shared else next(iter(ka))
        return {"relation": "same", "actor_a": k, "actor_b": k,
                "unit_a": ka.get(k, xa[0])["unit"], "unit_b": kb.get(k, xb[0])["unit"]}
    broad_a = all(k.startswith("class:") for k in ka)
    broad_b = all(k.startswith("class:") for k in kb)
    if level == "weak":
        rel = ("broad_same_class" if broad_a and broad_b else
               "broad_covers" if broad_a or broad_b else "sibling")
    else:
        rel = {"nested": "subtype", "different": "different"}[level]
    return {"relation": rel, "actor_a": actor_key(xa[0]), "actor_b": actor_key(xb[0]),
            "unit_a": xa[0]["unit"], "unit_b": xb[0]["unit"]}


def stratum_of(level, d, ra, rb):
    if level == "same":
        def match(rec, key, unit):
            for a in usable_actors(rec):
                if actor_key(a) == key and (unit is None or a["unit"] == unit):
                    return a.get("match")
            return None
        ms = {match(ra, d["actor_a"], d.get("unit_a")), match(rb, d["actor_b"], d.get("unit_b"))}
        return "same | shortened match" if "shortened" in ms else "same | both exact"
    if level == "nested":
        return "nested"
    if level == "weak":
        rel = d["relation"]
        if rel == "broad_covers":
            return "weak | broad covers mapped"
        if rel == "broad_same_class":
            return "weak | both broad, same class"
        return "weak | siblings"
    return ("different | same document" if ra.get("document") == rb.get("document")
            else "different | other documents")


# ---------------------------------------------------------------------------
# Statistics
# ---------------------------------------------------------------------------

def wilson(p, n, z=Z):
    if n <= 0:
        return 0.0, 1.0
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(max(0.0, p * (1 - p) / n + z * z / (4 * n * n))) / d
    return max(0.0, c - h), min(1.0, c + h)


def combine(parts):
    """parts: (N_h, k, n). Stratified proportion, Wilson on the effective n (capped)."""
    parts = [x for x in parts if x[0] > 0 and x[2] > 0]
    if not parts:
        return None
    total = sum(w for w, _, _ in parts)
    p = var = 0.0
    n_all = 0
    for w, k, n in parts:
        ph = k / n
        p += w / total * ph
        var += (w / total) ** 2 * max(0.0, 1 - n / w) * ph * (1 - ph) / n
        n_all += n
    n_eff = min(n_all, p * (1 - p) / var) if var > 0 else n_all
    lo, hi = wilson(p, n_eff)
    return p, lo, hi, n_eff


def fmt(res):
    if res is None:
        return "(no labels)"
    p, lo, hi, n_eff = res
    return f"{p:5.1%} [{lo:5.1%} - {hi:5.1%}]   n_eff {n_eff:5.1f}"


# ---------------------------------------------------------------------------
# Sampling
# ---------------------------------------------------------------------------

def sample(args):
    data = Path(args.data)
    out = Path(args.out) if args.out else data / "actor_pair_review.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)

    norm_path = Path(args.norms) if Path(args.norms).is_file() else data / args.norms
    score_path = Path(args.scores) if Path(args.scores).is_file() else data / args.scores
    for p, what in ((norm_path, "--norms"), (score_path, "--scores"), (Path(args.stage_a), "--stage-a")):
        if not p.is_file():
            raise SystemExit(f"{what}: {p} not found.")
    recs = {str(r["norm_id"]): r for r in read_jsonl(norm_path)}
    stage = json.loads(Path(args.stage_a).read_text(encoding="utf-8"))
    ids = [str(x) for x in stage["norm_ids"]]
    fmt_ = list(stage.get("pair_format") or ["pair_id", "i", "j"])
    ci, cj = fmt_.index("i"), fmt_.index("j")
    cp = fmt_.index("pair_id") if "pair_id" in fmt_ else None
    pairs = {(p[cp] if cp is not None else k): (p[ci], p[cj]) for k, p in enumerate(stage["pairs"])}
    scores = {r["pair_id"]: r for r in read_jsonl(score_path)}

    missing_ids = sorted(set(ids) - set(recs))
    if missing_ids:
        raise SystemExit(f"{len(missing_ids)} Stage A norms are not in {norm_path}, e.g. {missing_ids[:3]}.")
    use = {nid: bool(usable_actors(recs[nid])) for nid in ids}

    # --- population checks: pair ids are only valid with the Stage A file they came from
    outside = [pid for pid in scores if pid not in pairs]
    no_actor = [pid for pid in scores if pid in pairs
                and not (use[ids[pairs[pid][0]]] and use[ids[pairs[pid][1]]])]
    if outside or no_actor:
        raise SystemExit(
            f"actor_pair_scores.jsonl does not fit {args.stage_a}: {len(outside):,} pair ids are not "
            f"in it and {len(no_actor):,} decode to a norm without an actor. The scores were made "
            "with another Stage A file (pair ids depend on the order of its norm list). Re-run "
            "actor_scores.py with --stage-a pointing to this file.")
    bad = Counter(r["score"] for r in scores.values() if r["score"] not in LEVEL_OF_SCORE)
    if bad:
        raise SystemExit(f"Unexpected scores {dict(bad)}: this review assumes the default level "
                         "scores (1.0 / 0.7 / 0.2 / -0.25).")

    scorer = load_scorer(args, data, norm_path)

    # --- exact counts and the scored pairs with their strata
    counts = Counter()
    status_pairs = Counter()
    unknown_src = Counter()
    unknown_eunl = []
    pools = defaultdict(list)
    mismatch = 0
    for pid, (i, j) in pairs.items():
        a, b = ids[i], ids[j]
        if not (use[a] and use[b]):
            counts["no usable actor"] += 1
            for st in {recs[x]["status"] for x in (a, b) if not use[x]}:
                status_pairs[st] += 1
            continue
        sp = "-".join(sorted((source_of(a), source_of(b))))
        r = scores.get(pid)
        if r is None:
            counts["unknown"] += 1
            unknown_src[sp] += 1
            if sp == "EU-NL" and args.check_unknown:
                unknown_eunl.append(pid)
            continue
        level = LEVEL_OF_SCORE[r["score"]]
        counts[level] += 1
        d = None
        if scorer is not None:
            full = scorer.norm_pair(a, b)
            if abs(full["score"] - r["score"]) < 1e-9 and full.get("relation"):
                d = {k: full.get(k) for k in ("relation", "actor_a", "actor_b", "unit_a", "unit_b")}
                d["source"] = "scorer"
            else:
                mismatch += 1
        if d is None:
            d = approx_detail(level, recs[a], recs[b])
            d["source"] = "approx"
        st = stratum_of(level, d, recs[a], recs[b])
        pools[st].append({"pair_id": pid, "a": a, "b": b, "score": r["score"], "level": level,
                          "direction": r.get("direction") or "", "sources": sp, **d})
    if scorer is not None and mismatch:
        share = mismatch / max(1, sum(len(v) for v in pools.values()))
        msg = (f"ActorScorer disagrees with the score file on {mismatch:,} pairs ({share:.1%}). "
               "The relations, inventory or curated file differ from the scoring run, or the "
               "scores are stale.")
        if share > args.max_mismatch:
            raise SystemExit(msg + f" That is more than --max-mismatch ({args.max_mismatch:.0%}): "
                             "re-run actor_scores.py, or raise --max-mismatch to use the "
                             "approximate relation for those pairs.")
        print("Warning: " + msg + " Those pairs use the approximate relation ('approx' in the key).")

    # dominant same IDs, for a descriptive breakdown
    same_ids = Counter(p["actor_a"] for s in ("same | both exact", "same | shortened match")
                       for p in pools[s])
    top_same = {k for k, _ in same_ids.most_common(args.top_same)}

    alloc = {s: n for s, _, n, _ in STRATA}
    for item in args.alloc or []:
        name, _, n = item.rpartition("=")
        if name.strip() not in alloc:
            raise SystemExit(f"--alloc: unknown stratum '{name.strip()}'. Strata: {list(alloc)}")
        alloc[name.strip()] = int(n)
    N = {s: len(pools[s]) for s, *_ in STRATA}

    # --- draw: smallest strata first, at most --max-uses pairs per norm
    rng = random.Random(args.seed)
    uses = Counter()
    rows = []
    for s in sorted(alloc, key=lambda s: N[s]):
        pool = sorted(pools[s], key=lambda p: p["pair_id"])
        rng.shuffle(pool)
        taken = 0
        for p in pool:
            if taken == alloc[s]:
                break
            if uses[p["a"]] >= args.max_uses or uses[p["b"]] >= args.max_uses:
                continue
            uses[p["a"]] += 1
            uses[p["b"]] += 1
            taken += 1
            rows.append(dict(p, stratum=s, kind="sample",
                             top_same="yes" if p["actor_a"] in top_same else "no"))
        if taken < alloc[s]:
            print(f"Warning: '{s}' yielded {taken} pairs (target {alloc[s]}).")
    if args.check_unknown:
        for pid in rng.sample(sorted(unknown_eunl), min(args.check_unknown, len(unknown_eunl))):
            i, j = pairs[pid]
            a, b = ids[i], ids[j]
            xa, xb = usable_actors(recs[a]), usable_actors(recs[b])
            rows.append({"pair_id": pid, "a": a, "b": b, "score": 0.0, "level": "unknown",
                         "direction": "", "sources": "EU-NL", "relation": "unknown",
                         "actor_a": actor_key(xa[0]), "actor_b": actor_key(xb[0]),
                         "unit_a": xa[0]["unit"], "unit_b": xb[0]["unit"], "source": "map",
                         "stratum": CHECK_UNKNOWN, "kind": "check", "top_same": "no"})
    rng.shuffle(rows)   # order must not reveal the stratum

    texts = {}
    if args.norm_text:
        for r in read_records(args.norm_text):
            texts[str(_get(r, args.id_field))] = str(_get(r, args.text_field) or "")[:args.text_chars]

    info = dict(N=N, alloc=alloc, taken=Counter(r["stratum"] for r in rows), counts=counts,
                status_pairs=status_pairs, unknown_src=unknown_src, total=len(pairs),
                scorer=scorer is not None, mismatch=mismatch, args=args)
    write_workbook(out, rows, recs, texts, info)

    print(f"Wrote {len(rows)} pairs -> {out}")
    print(f"Stage A pairs {len(pairs):,}: no usable actor {counts['no usable actor']:,}, "
          f"unknown {counts['unknown']:,}, scored {sum(N.values()):,}")
    print(f"Relations from: {'ActorScorer' if scorer else 'the actor map (approx)'}")
    print(f"{'stratum':34} {'N_h':>8} {'n':>4}")
    for s, *_ in STRATA:
        print(f"{s:34} {N[s]:>8,} {info['taken'][s]:>4}")
    if args.check_unknown:
        print(f"{CHECK_UNKNOWN:34} {unknown_src['EU-NL']:>8,} {info['taken'][CHECK_UNKNOWN]:>4}")
    if not texts:
        print("No --norm-text given: the norm text columns are empty.")


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def header(ws, row, ncols, col0=1):
    for c in range(col0, col0 + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def put(ws, row, col, value, wrap=False, bold=False, italic=False, numfmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, bold=bold, italic=italic)
    c.alignment = Alignment(vertical="top", wrap_text=wrap)
    if numfmt:
        c.number_format = numfmt
    return c


def expected_halfwidth(info, p=0.5):
    N, alloc = info["N"], info["alloc"]
    names = [s for s, *_ in STRATA if alloc[s] and N[s]]
    tot = sum(N[s] for s in names)
    v = sum((N[s] / tot) ** 2 * max(0.0, 1 - alloc[s] / N[s]) * p * (1 - p) / alloc[s] for s in names)
    return 100 * Z * math.sqrt(v)


def design_note(info):
    c, N, a = info["counts"], info["N"], info["args"]
    scored = sum(N.values())
    hw50, hw80 = expected_halfwidth(info, 0.5), expected_halfwidth(info, 0.8)
    return (
        "What is validated. For each Stage A pair the actor score compares the actors of the two "
        "norms: same (1.0, same actor ID), nested (0.7, one actor is a kind or part of the other), "
        "weak (0.2, two kinds of one actor, or a broad actor that may include the other), different "
        "(-0.25, two unrelated actors from a shared document). The labeller answers the same "
        "question from the norm texts: do the two norms address the same actor?\n\n"
        f"Population. {info['total']:,} Stage A pairs. {scored:,} have a score and are sampled. "
        f"{c['no usable actor']:,} involve a norm without a usable actor and {c['unknown']:,} are "
        "'unknown' (both norms have an actor, but no relation is known); these score 0, make no "
        "claim, and are exact counts (table below). Related actors among them cannot be measured "
        "here. Note that an EU and a Dutch actor ID are never related, so EU-Dutch pairs are only "
        "ever weak (through broad classes) or unknown.\n\n"
        "Why these strata. The strata are the score levels, split by the relation behind each "
        "pair, because each relation can fail in its own way: 'same' by exact or by shortened "
        "match; 'weak' as siblings, a broad class covering a mapped actor, or two broad actors; "
        "'different' within one document (where a missed relation is most likely) or across "
        "documents. Nested is small and kept whole.\n\n"
        "Why these sizes. Weak and different hold most scored pairs, so they get most judgements; "
        "same and nested are small but make the strongest claims, so they get floors. A norm "
        f"appears in at most {a.max_uses} sampled pairs. With this allocation the weighted "
        f"agreement over all scored pairs is known to about +/-{hw80:.0f} points at 80% agreement "
        f"(+/-{hw50:.0f} at 50%). Overall figures are weighted by N_h (the 'review' step), never "
        "by row count.\n\n"
        "Relations. " + (("Taken from ActorScorer with the scoring run's relations and inventory"
                          + (f"; {info['mismatch']:,} pairs where it disagreed with the score file "
                             "use the approximate relation." if info["mismatch"] else "."))
                         if info["scorer"] else
                         "ActorScorer was not available, so the relation behind weak pairs was "
                         "derived from the actor map (marked 'approx' in the key); the weak "
                         "strata are approximate.")
    )


def write_workbook(path, rows, recs, texts, info):
    N, a = info["N"], info["args"]
    wb = Workbook()

    # ---------------- Instructions ----------------
    g = wb.active
    g.title = "Instructions"
    lines = [
        ("How to judge", True),
        ("Each row of the Review sheet is a pair of norms, A and B. Judge only their actors: who "
         "is addressed by each norm. Use the norm texts.", False),
        ("Actor relation?", True),
        ("  same       both norms address the same actor (other wording counts).", False),
        ("  nested     one actor is a kind or a part of the other (e.g. 'aanbieder van een AI-systeem "
         "met een hoog risico' in 'aanbieder').", False),
        ("  related    different but closely related actors (two kinds of one actor), or a broad "
         "actor that may include the other.", False),
        ("  different  different actors.", False),
        ("  unclear    cannot tell from the texts (not counted).", False),
        ("Narrower actor?   only when 'same' or 'nested': A narrower / B narrower / equal / can't say.", False),
        ("Actor units show how each actor phrase was split; judge the actor as a whole. When a norm "
         "has several actors, judge the closest pair.", False),
        ("Scores are in the hidden 'key' sheet: leave it alone until you are done.", False),
        ("When finished, save and run: python actor_pair_review.py review --xlsx <this file>", False),
        ("", False),
        ("Example (format only; not part of the sample)", True),
        ("Actor A: de aanbieder van een AI-systeem met een hoog risico | Actor B: aanbieders", False),
        ("Actor relation? nested | Narrower actor? A narrower", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        put(g, i, 1, text, bold=bold, wrap=True)
    g.column_dimensions["A"].width = 110

    # ---------------- Review (blind) ----------------
    ws = wb.create_sheet("Review")
    cols = [("#", 5), ("Norm A", 22), ("Actor A", 34), ("Actor units A", 26)]
    if texts:
        cols += [("Text A", 55)]
    cols += [("Norm B", 22), ("Actor B", 34), ("Actor units B", 26)]
    if texts:
        cols += [("Text B", 55)]
    cols += [("Actor relation?", 14), ("Narrower actor?", 14), ("Note", 30)]
    names = [c for c, _ in cols]
    for j, (name, width) in enumerate(cols, 1):
        ws.cell(row=1, column=j, value=name)
        ws.column_dimensions[get_column_letter(j)].width = width
    header(ws, 1, len(cols))
    for r, p in enumerate(rows, 2):
        ra, rb = recs[p["a"]], recs[p["b"]]
        values = {"#": r - 1, "Norm A": p["a"], "Actor A": ra.get("actor") or "",
                  "Actor units A": " | ".join(x["unit"] for x in usable_actors(ra)),
                  "Text A": texts.get(p["a"], ""),
                  "Norm B": p["b"], "Actor B": rb.get("actor") or "",
                  "Actor units B": " | ".join(x["unit"] for x in usable_actors(rb)),
                  "Text B": texts.get(p["b"], "")}
        for j, name in enumerate(names, 1):
            c = put(ws, r, j, values.get(name), wrap=name != "#")
            c.border = Border(bottom=THIN)
            if name in ("Actor relation?", "Narrower actor?", "Note"):
                c.fill = INPUT_FILL
    last = len(rows) + 1
    lcol = get_column_letter(names.index("Actor relation?") + 1)
    ncol = get_column_letter(names.index("Narrower actor?") + 1)
    for col, options in ((lcol, LABELS), (ncol, NARROWER)):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    ws.freeze_panes = "B2"

    # ---------------- key (hidden) ----------------
    ks = wb.create_sheet("key")
    kh = ["#", "stratum", "level", "score", "relation", "detail_from", "actor_a", "actor_b",
          "unit_a", "unit_b", "direction", "sources", "top_same_id", "pair_id"]
    ks.append(kh)
    for i, p in enumerate(rows, 1):
        ks.append([i, p["stratum"], p["level"], p["score"], p["relation"], p["source"],
                   p["actor_a"], p["actor_b"], p["unit_a"], p["unit_b"], p["direction"],
                   p["sources"], p["top_same"], p["pair_id"]])
    for row in ks.iter_rows():
        for c in row:
            c.font = Font(name=FONT)
    ks.sheet_state = "hidden"

    # ---------------- Design ----------------
    ds = wb.create_sheet("Design")
    put(ds, 1, 1, "Sample design", bold=True)
    ds.merge_cells("A2:F2")
    put(ds, 2, 1, design_note(info), wrap=True)
    ds.row_dimensions[2].height = 400
    th = ["stratum", "level", "population N_h", "share of scored pairs", "n sampled",
          "why this stratum and this n"]
    hr = 4
    for j, h in enumerate(th, 1):
        ds.cell(row=hr, column=j, value=h)
    header(ds, hr, len(th))
    first = hr + 1
    lastrow = hr + len(STRATA)
    for k, (s, lvl, _, why) in enumerate(STRATA):
        r = first + k
        put(ds, r, 1, s)
        put(ds, r, 2, lvl)
        put(ds, r, 3, N[s], numfmt="#,##0")
        put(ds, r, 4, f"=IFERROR(C{r}/SUM($C${first}:$C${lastrow}),0)", numfmt="0.0%")
        put(ds, r, 5, info["taken"][s])
        put(ds, r, 6, why, wrap=True)
    tr = lastrow + 1
    put(ds, tr, 1, "total", bold=True)
    put(ds, tr, 3, f"=SUM(C{first}:C{lastrow})", bold=True, numfmt="#,##0")
    put(ds, tr, 5, f"=SUM(E{first}:E{lastrow})", bold=True)

    r = tr + 2
    put(ds, r, 1, "Exact counts (not sampled)", bold=True)
    for j, h in enumerate(["pairs", "count", "share of Stage A"], 1):
        ds.cell(row=r + 1, column=j, value=h)
    header(ds, r + 1, 3)
    c, sp = info["counts"], info["status_pairs"]
    exact = [("Stage A pairs", info["total"]),
             ("scored (sampled above)", sum(N.values())),
             ("a norm without a usable actor", c["no usable actor"]),
             ("   ... with a norm of status not_actor", sp["not_actor"]),
             ("   ... with a norm of status missing", sp["missing"]),
             ("   ... with a norm of status pronoun", sp["pronoun"]),
             ("   ... with a norm of status no_units", sp["no_units"]),
             ("unknown: both have an actor, no relation known", c["unknown"]),
             ("   ... EU-EU", info["unknown_src"]["EU-EU"]),
             ("   ... EU-NL", info["unknown_src"]["EU-NL"]),
             ("   ... NL-NL", info["unknown_src"]["NL-NL"])]
    for k, (label, v) in enumerate(exact):
        rr = r + 2 + k
        put(ds, rr, 1, label)
        put(ds, rr, 2, v, numfmt="#,##0")
        put(ds, rr, 3, f"=B{rr}/$B${r + 2}", numfmt="0.0%")
    rr = r + 2 + len(exact)
    put(ds, rr, 1, "A pair can have two norms without an actor of different status, so the status "
                   "lines overlap.", italic=True)
    put(ds, rr + 1, 1, f"Settings: max uses per norm {a.max_uses}, seed {a.seed}, "
                       f"check unknown {a.check_unknown}, relations from "
                       f"{'ActorScorer' if info['scorer'] else 'actor map (approx)'}", italic=True)
    for col, w in zip("ABCDEF", [44, 12, 14, 12, 10, 70]):
        ds.column_dimensions[col].width = w

    # ---------------- Summary (live, unweighted per stratum) ----------------
    sm = wb.create_sheet("Summary")
    sh = ["stratum", "expected label", "population N_h", "rows", "judged"] + LABELS + \
         ["agreement", "agree 95% low", "agree 95% high", "share same/nested"]
    sm.append(sh)
    header(sm, 1, len(sh))
    sm.cell(row=1, column=17, value="z")
    sm.cell(row=2, column=17, value=Z)
    srng, lrng = f"key!$B$2:$B${last}", f"Review!${lcol}$2:${lcol}${last}"
    names_sum = [s for s, *_ in STRATA] + ([CHECK_UNKNOWN] if a.check_unknown else [])
    for i, s in enumerate(names_sum, start=2):
        exp = EXPECTED[STRATUM_LEVEL[s]] if s in STRATUM_LEVEL else "different"
        put(sm, i, 1, s)
        put(sm, i, 2, exp)
        put(sm, i, 3, N.get(s, info["unknown_src"]["EU-NL"] if s == CHECK_UNKNOWN else 0), numfmt="#,##0")
        put(sm, i, 4, f"=COUNTIFS({srng},A{i})")
        put(sm, i, 5, f'=COUNTIFS({srng},A{i},{lrng},"<>")')
        for j, lab in enumerate(LABELS, start=6):
            put(sm, i, j, f'=COUNTIFS({srng},A{i},{lrng},"{lab}")')
        n = f"(E{i}-J{i})"
        put(sm, i, 11, f'=IFERROR(COUNTIFS({srng},A{i},{lrng},B{i})/{n},"")', numfmt="0%")
        core = f"(K{i}+$Q$2^2/(2*{n}))/(1+$Q$2^2/{n})"
        half = f"$Q$2*SQRT(K{i}*(1-K{i})/{n}+$Q$2^2/(4*{n}^2))/(1+$Q$2^2/{n})"
        put(sm, i, 12, f'=IFERROR(MAX(0,{core}-{half}),"")', numfmt="0%")
        put(sm, i, 13, f'=IFERROR(MIN(1,{core}+{half}),"")', numfmt="0%")
        put(sm, i, 14, f'=IFERROR((F{i}+G{i})/{n},"")', numfmt="0%")
    nr = len(names_sum) + 3
    for k, t in enumerate([
            "Per stratum, unweighted. 'agreement' = share of judged pairs whose label is the one the "
            "score level claims (same, nested, related, different); unclear is excluded.",
            "For the unknown check, 'agreement' counts 'different' as agreeing (no overlap missed).",
            "Weighted figures, the overlap decision, the direction check and error kinds: "
            "python actor_pair_review.py review --xlsx <this file>"]):
        put(sm, nr + k, 1, t, italic=True)
    for col, w in zip("ABCDEFGHIJKLMN", [34, 12, 13, 6, 7, 7, 7, 8, 9, 8, 10, 12, 12, 13]):
        sm.column_dimensions[col].width = w

    wb.active = wb.sheetnames.index("Review")
    wb.save(path)


# ---------------------------------------------------------------------------
# Review
# ---------------------------------------------------------------------------

ERROR_KIND = {
    ("same", "nested"): "sub-role taken for the same actor",
    ("same", "related"): "related actors merged into one ID",
    ("same", "different"): "wrong actor ID (mapping error)",
    ("nested", "same"): "same actor scored as nested (equivalence missing)",
    ("nested", "related"): "nesting relation too strong",
    ("nested", "different"): "wrong nesting relation",
    ("weak", "same"): "overlap scored only weak (relation missing)",
    ("weak", "nested"): "overlap scored only weak (relation missing)",
    ("weak", "different"): "weak link wrong (class too broad / wrong sibling)",
    ("different", "same"): "same actor under two IDs (equivalence missing)",
    ("different", "nested"): "nested relation missing",
    ("different", "related"): "weak relation missing",
}


def review(args):
    wb = load_workbook(args.xlsx, data_only=True)
    rv = wb["Review"]
    names = [c.value for c in rv[1]]
    lc, nc = names.index("Actor relation?"), names.index("Narrower actor?")
    lab = {}
    for row in rv.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            lab[int(row[0])] = (str(row[lc]).strip().lower() if row[lc] else None,
                                str(row[nc]).strip() if row[nc] else None)
    kh = [c.value for c in wb["key"][1]]
    key = {int(r[0]): dict(zip(kh, r)) for r in wb["key"].iter_rows(min_row=2, values_only=True)
           if r[0] is not None}
    N = {}
    for row in wb["Design"].iter_rows(values_only=True):
        if row and row[0] in STRATUM_LEVEL and isinstance(row[2], (int, float)):
            N[row[0]] = int(row[2])

    judged = {i: v for i, v in lab.items() if v[0] in LABELS}
    print(f"Rows {len(lab)}   judged {len(judged)}   unclear "
          f"{sum(v[0] == 'unclear' for v in judged.values())}   "
          f"blank {sum(1 for v in lab.values() if not v[0])}\n")
    if not judged:
        print("Nothing to review yet - fill in 'Actor relation?' first.")
        return

    def rows_of(stratum):
        return [(key[i], v) for i, v in judged.items()
                if v[0] != "unclear" and key[i]["stratum"] == stratum]

    def part(stratum, ok):
        xs = rows_of(stratum)
        return (N.get(stratum, 0), sum(ok(k, v) for k, v in xs), len(xs))

    agree = lambda k, v: v[0] == EXPECTED[k["level"]]
    overlap_ok = lambda k, v: (v[0] in OVERLAP) == (k["level"] in OVERLAP)
    is_overlap = lambda k, v: v[0] in OVERLAP
    strata = [s for s, *_ in STRATA]

    print("Per stratum (95% Wilson):")
    print(f"  {'stratum':34} {'N_h':>8} {'n':>3}  {'agrees with level':22} {'labelled same/nested':22}")
    for s in strata:
        _, k, n = part(s, agree)
        _, ko, _ = part(s, is_overlap)
        if not n:
            print(f"  {s:34} {N.get(s, 0):>8,} {0:>3}  (no labels)")
            continue
        lo, hi = wilson(k / n, n)
        lo2, hi2 = wilson(ko / n, n)
        print(f"  {s:34} {N[s]:>8,} {n:>3}  {k/n:4.0%} [{lo:3.0%}-{hi:4.0%}]       "
              f"{ko/n:4.0%} [{lo2:3.0%}-{hi2:4.0%}]")

    print("\nWeighted by N_h over all scored pairs:")
    print(f"  {'agreement with the score level':40} {fmt(combine([part(s, agree) for s in strata]))}")
    print(f"  {'overlap decision (score >= 0.5) correct':40} "
          f"{fmt(combine([part(s, overlap_ok) for s in strata]))}")
    hi_strata = [s for s in strata if STRATUM_LEVEL[s] in OVERLAP]
    print(f"  {'precision of the overlap (same/nested)':40} "
          f"{fmt(combine([part(s, is_overlap) for s in hi_strata]))}")
    for lvl in ("same", "nested", "weak", "different"):
        ss = [s for s in strata if STRATUM_LEVEL[s] == lvl]
        print(f"  {'agreement, level ' + lvl:40} {fmt(combine([part(s, agree) for s in ss]))}")

    for lvl in ("weak", "different"):
        ss = [s for s in strata if STRATUM_LEVEL[s] == lvl]
        res = combine([part(s, is_overlap) for s in ss])
        if res:
            tot = sum(N[s] for s in ss)
            print(f"\nMissed overlaps among {lvl} pairs (labelled same or nested):")
            print(f"  share  {fmt(res)}")
            print(f"  pairs  ~{res[0] * tot:,.0f} of {tot:,} [{res[1] * tot:,.0f} - {res[2] * tot:,.0f}]")

    # confusion of level x label, weighted to population counts
    print("\nScore level x label (sampled count; weighted estimate of pairs):")
    labs = ["same", "nested", "related", "different"]
    print("  " + f"{'level':10}" + "".join(f"{l:>20}" for l in labs))
    kinds = Counter()
    for lvl in ("same", "nested", "weak", "different"):
        cells = []
        for l in labs:
            cnt, est = 0, 0.0
            for s in [s for s in strata if STRATUM_LEVEL[s] == lvl]:
                xs = rows_of(s)
                k = sum(v[0] == l for _, v in xs)
                cnt += k
                if xs:
                    est += N[s] * k / len(xs)
            cells.append(f"{cnt:>4} (~{est:>10,.0f})")
            if (lvl, l) in ERROR_KIND and cnt:
                kinds[ERROR_KIND[(lvl, l)]] += est
        print("  " + f"{lvl:10}" + "".join(f"{c:>20}" for c in cells))
    if kinds:
        print("\nError kinds (weighted estimate of pairs):")
        for kname, est in kinds.most_common():
            print(f"  ~{est:>10,.0f}  {kname}")

    # direction
    nested = [(key[i], v) for i, v in judged.items() if key[i]["level"] == "nested"
              and v[0] == "nested" and v[1] in FLAG_ANSWER.values()
              and key[i]["direction"] in FLAG_ANSWER]
    print("\nDirection flag (nested pairs labelled nested, answer given):")
    if nested:
        ok = sum(v[1] == FLAG_ANSWER[k["direction"]] for k, v in nested)
        lo, hi = wilson(ok / len(nested), len(nested))
        print(f"  agreement {ok}/{len(nested)} = {ok/len(nested):.0%} [{lo:.0%} - {hi:.0%}]")
        print("  flag        " + "".join(f"{b:>13}" for b in NARROWER[:2]))
        tab = Counter((k["direction"], v[1]) for k, v in nested)
        for f in FLAG_ANSWER:
            print(f"  {f:12}" + "".join(f"{tab[(f, b)]:>13}" for b in NARROWER[:2]))
    else:
        print("  none yet")
    same_eq = [(k, v) for i, v in judged.items() if (k := key[i])["level"] == "same"
               and v[0] == "same" and v[1] and v[1] != "can't say"]
    if same_eq:
        eq = sum(v[1] == "equal" for _, v in same_eq)
        print(f"  same pairs labelled same, answered 'equal': {eq}/{len(same_eq)}")

    # descriptive breakdowns
    print("\nDescriptive (unweighted; groups too small to compare):")
    for title, field, vals in (("source pair", "sources", ("EU-EU", "EU-NL", "NL-NL")),
                               ("same pairs, dominant IDs", "top_same_id", ("yes", "no"))):
        for val in vals:
            xs = [(k, v) for i, v in judged.items() if v[0] != "unclear"
                  and (k := key[i])["stratum"] in strata and k[field] == val
                  and (field != "top_same_id" or k["level"] == "same")]
            if xs:
                ok = sum(agree(k, v) for k, v in xs)
                lo, hi = wilson(ok / len(xs), len(xs))
                print(f"  {title:26} {val:6} agreement {ok}/{len(xs)} = {ok/len(xs):.0%} [{lo:.0%} - {hi:.0%}]")
    approx = sum(1 for k in key.values() if k["detail_from"] == "approx")
    if approx:
        print(f"\nNote: {approx} sampled pairs have an approximate relation (no ActorScorer at sampling).")

    chk = [(key[i], v) for i, v in judged.items() if key[i]["stratum"] == CHECK_UNKNOWN and v[0] != "unclear"]
    if chk:
        miss = sum(v[0] in OVERLAP for _, v in chk)
        lo, hi = wilson(miss / len(chk), len(chk))
        print(f"\nCheck, unknown EU-NL pairs labelled same/nested: {miss}/{len(chk)} = "
              f"{miss/len(chk):.0%} [{lo:.0%} - {hi:.0%}]")

    wrong = [(key[i], v) for i, v in judged.items() if key[i]["stratum"] in strata
             and v[0] != "unclear" and not agree(key[i], v)]
    if wrong:
        print("\nDisagreements:")
        for k, v in sorted(wrong, key=lambda x: (x[0]["level"], x[0]["#"])):
            print(f"  #{k['#']:<4} {k['level']:9} {k['relation']:16} {k['actor_a']!s:40} | "
                  f"{k['actor_b']!s:40} -> {v[0]}")


# ---------------------------------------------------------------------------

def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample", help="Write the labelling workbook.")
    s.add_argument("--data", default="data/conflict_all/actor", help="folder of the actor step's output")
    s.add_argument("--norms", default="norm_actor_map.jsonl", help="path, or file name in --data")
    s.add_argument("--scores", default="actor_pair_scores.jsonl", help="path, or file name in --data")
    s.add_argument("--stage-a", default="data/conflict_all/stageA_pairs.json",
                   help="the Stage A file the scores were made with")
    s.add_argument("--relations", default="actor_relations.jsonl", help="in --data, for ActorScorer")
    s.add_argument("--inventory", default="actor_inventory.jsonl", help="in --data, for ActorScorer")
    s.add_argument("--curated", default=None,
                   help="curated relations (default: as recorded in actor_scores_config.json)")
    s.add_argument("--scorer", default=None, help="path to actor_scores.py (default: next to this script)")
    s.add_argument("--max-mismatch", type=float, default=0.01,
                   help="stop if ActorScorer disagrees with the score file on more than this share "
                        "of scored pairs (default 0.01)")
    s.add_argument("--norm-text", default=None, help="norm records with the text (JSON array or JSONL)")
    s.add_argument("--id-field", default="norm_id")
    s.add_argument("--text-field", default="text.paragraph")
    s.add_argument("--text-chars", type=int, default=1500)
    s.add_argument("--alloc", action="append", metavar="STRATUM=N",
                   help='override a stratum size, e.g. --alloc "nested=25"')
    s.add_argument("--max-uses", type=int, default=2)
    s.add_argument("--top-same", type=int, default=3, help="dominant same IDs for the breakdown")
    s.add_argument("--check-unknown", type=int, default=0,
                   help="also draw N unknown EU-NL pairs (reported separately)")
    s.add_argument("--seed", type=int, default=42)
    s.add_argument("--out", default=None)
    r = sub.add_parser("review", help="Report estimates from a labelled workbook.")
    r.add_argument("--xlsx", default="data/conflict/actor/actor_pair_review.xlsx")
    args = p.parse_args()
    sample(args) if args.cmd == "sample" else review(args)


if __name__ == "__main__":
    main()