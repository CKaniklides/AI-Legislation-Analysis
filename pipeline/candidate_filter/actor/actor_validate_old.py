"""
Actor filter, validation: a review workbook for a human, and its scores.

  sample  draws a stratified random sample of mapped actor units and writes a
          workbook for a reviewer (Instructions, Review, Scores, Inventory).
  score   reads the filled workbook and reports the validation scores, with
          95% confidence intervals and the kinds of errors found.

The Scores sheet already computes the main scores with formulas while the
reviewer works; 'score' adds intervals and an error breakdown.

What is sampled
  One case is one distinct actor unit in one document ('Aanbieders van
  AI-systemen met een hoog risico' in 32024R1689), shown with one random norm
  that uses it. Cases are drawn per match type (the strata):

      exact 14, shortened 10, fallback 8, not_actor 8     (40 cases)

  Exact matches are most of the data but rarely wrong, so the riskier strata
  are oversampled. Within a stratum, cases are spread over documents and over
  actor IDs (or fallback classes, or head words): one case per group, round by
  round, so the sample is not dominated by 'Commissie'. The overall score is
  weighted back to the real share of each stratum (by norm uses).

  40 cases give about +/-11 points (95%) on an accuracy near 85%: enough to
  tell whether the mapping is above the 75% target, in about 40 minutes of
  review. --size (20-50) changes it; the strata keep their proportions.

Inputs
  --dir      folder with unit_actor_map.jsonl, norm_actor_map.jsonl and
             actor_inventory.jsonl (actor_map.py, actor_inventory.py)
  --records  optional norm records with the norm text, shown to the reviewer
             (--text-field, dotted path, default 'norm.text'; --id-field)
Outputs
  sample: actor_validation.xlsx (or --out)
  score:  scores printed, and <workbook>.scores.json

Usage
  python actor_validate.py sample --dir data/conflict/actor --records data/norms.jsonl
  python actor_validate.py score  --file data/conflict/actor/actor_validation.xlsx
"""

from __future__ import annotations

import argparse
import json
import math
import random
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

STRATA = ["exact", "shortened", "fallback", "not_actor"]
QUOTA_40 = {"exact": 14, "shortened": 10, "fallback": 8, "not_actor": 8}
VERDICTS = ["yes", "no", "unsure"]

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="D9E1F2")
INPUT_FILL = PatternFill("solid", start_color="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Review sheet columns (letter, header, width)
COLS = [
    ("A", "case", 6), ("B", "document", 12), ("C", "norm_id", 18),
    ("D", "raw actor (as extracted)", 34), ("E", "actor unit (cleaned)", 30),
    ("F", "norm text", 60), ("G", "system: match", 11), ("H", "system: actor ID", 34),
    ("I", "system: label", 26), ("J", "system: fallback class", 12),
    ("K", "REVIEW: system correct?", 13), ("L", "REVIEW: correct answer (if no)", 34),
    ("M", "REVIEW: notes", 30),
]
FIRST = 2  # first data row of the Review sheet


# --------------------------------------------------------------------------
# Sampling
# --------------------------------------------------------------------------

def quotas(size: int) -> dict[str, int]:
    size = max(20, min(50, size))
    q = {s: max(2, round(QUOTA_40[s] * size / 40)) for s in STRATA}
    q["exact"] += size - sum(q.values())
    return q


def group_of(u: dict) -> tuple:
    return (u["document"], u.get("actor_id") or u.get("fallback_class") or u.get("head") or "")


def draw(units: list[dict], quota: dict[str, int], rng: random.Random) -> list[dict]:
    """Per stratum: one random unit per (document, ID/class/head) group, round
    by round, until the quota is filled."""
    picked = []
    for s in STRATA:
        groups = defaultdict(list)
        for u in units:
            if u["match"] == s:
                groups[group_of(u)].append(u)
        order = list(groups)
        rng.shuffle(order)
        # alternate documents in the group order
        by_doc = defaultdict(list)
        for g in order:
            by_doc[g[0]].append(g)
        order = [g for tup in _zip_longest(*by_doc.values()) for g in tup if g]
        for g in order:
            rng.shuffle(groups[g])
        n, take = quota[s], []
        while len(take) < n and any(groups[g] for g in order):
            for g in order:
                if groups[g] and len(take) < n:
                    take.append(groups[g].pop())
        picked += take
    return picked


def _zip_longest(*lists):
    for i in range(max((len(l) for l in lists), default=0)):
        yield tuple(l[i] if i < len(l) else None for l in lists)


def _get(obj, path: str):
    for k in path.split("."):
        if not isinstance(obj, dict):
            return None
        obj = obj.get(k)
    return obj


def load_texts(path: Path | None, id_field: str, text_field: str) -> dict[str, str]:
    if not path:
        return {}
    raw = path.read_text(encoding="utf-8").strip()
    records = json.loads(raw) if raw.startswith("[") else \
        [json.loads(l) for l in raw.splitlines() if l.strip()]
    return {str(_get(r, id_field)): str(_get(r, text_field) or "") for r in records}


# --------------------------------------------------------------------------
# Workbook
# --------------------------------------------------------------------------

def _style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = BORDER


def write_instructions(wb, size: int) -> None:
    ws = wb.active
    ws.title = "Instructions"
    lines = [
        ("Actor mapping review", True),
        ("", False),
        ("Each row of the Review sheet is one actor of a norm and the decision the system made "
         "for it. Judge the decision, using the norm text where it helps.", False),
        ("", False),
        ("What the system decides (column G)", True),
        ("exact / shortened: the actor is the inventory entry in column H. 'shortened' means "
         "words were dropped to find it (e.g. 'bevoegde toezichthoudende autoriteit' -> "
         "toezichthoudende autoriteit).", False),
        ("fallback: the unit is an actor, but none of the entries fits; only a broad class "
         "is given (column J).", False),
        ("not_actor: the unit is not an actor (a document, decision, system, ...).", False),
        ("", False),
        ("What to fill in (yellow cells only)", True),
        ("K  system correct?  yes / no / unsure.", False),
        ("     exact/shortened: yes if the actor in the norm is this entry (a sub-role of it "
         "also counts as yes, e.g. 'betrokken aanbieder' -> aanbieder).", False),
        ("     fallback: yes if it is an actor and no entry on the Inventory sheet fits it.", False),
        ("     not_actor: yes if the unit is indeed not an actor.", False),
        ("L  correct answer, only when K is 'no': an actor ID from the Inventory sheet, "
         "'fallback' (an actor without a fitting entry) or 'not_actor'.", False),
        ("M  notes: anything useful, e.g. 'unit is cut wrongly by the cleaning step'.", False),
        ("Only entries valid for the row's document may be chosen: a document's defined terms "
         "apply only to its own norms (Inventory sheet, column 'valid for').", False),
        ("", False),
        ("Example row (format only)", True),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=12 if i == 1 else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r = len(lines) + 1
    for j, (_, h, _) in enumerate(COLS, start=1):
        ws.cell(row=r, column=j, value=h)
    _style_header(ws, r, len(COLS))
    example = ["0", "32016R0679", "32016R0679:57:1:3", "De bevoegde toezichthoudende autoriteit",
               "bevoegde toezichthoudende autoriteit", "(norm text)", "shortened",
               "32016R0679:toezichthoudende_autoriteit", "toezichthoudende autoriteit", "",
               "yes", "", ""]
    for j, v in enumerate(example, start=1):
        c = ws.cell(row=r + 1, column=j, value=v)
        c.font = Font(name=FONT, italic=True, color="808080")
        c.border = BORDER
    ws.cell(row=r + 3, column=1,
            value=f"The sample has {size} cases. Scores update on the Scores sheet as you fill in "
                  "column K and L.").font = Font(name=FONT, italic=True)
    ws.column_dimensions["A"].width = 120
    for letter, _, width in COLS[1:]:
        ws.column_dimensions[letter].width = min(width, 20)


def write_inventory(wb, entries: list[dict]) -> int:
    ws = wb.create_sheet("Inventory")
    heads = ["actor ID", "label", "aliases", "kind", "valid for", "class"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=1, column=j, value=h)
    _style_header(ws, 1, len(heads))
    rows = [(e["id"], e["label"], "; ".join(e.get("aliases", [])), e["kind"],
             ", ".join(e["documents"]) if e["kind"] == "defined" else
             (e["scope"] if isinstance(e["scope"], str) else ", ".join(e["scope"]))
             + (f" (not in {', '.join(e['excluded'])})" if e.get("excluded") else ""),
             e.get("class")) for e in entries]
    rows += [("fallback", "an actor without a fitting entry", "", "", "all", ""),
             ("not_actor", "not an actor", "", "", "all", "")]
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v).font = Font(name=FONT)
    for letter, w in zip("ABCDEF", (44, 40, 40, 9, 34, 12)):
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A2"
    return len(rows) + 1  # last row


def write_review(wb, cases: list[dict], last_inv_row: int) -> int:
    ws = wb.create_sheet("Review")
    for j, (letter, h, w) in enumerate(COLS, start=1):
        ws.cell(row=1, column=j, value=h)
        ws.column_dimensions[letter].width = w
    _style_header(ws, 1, len(COLS))
    for i, c in enumerate(cases):
        r = FIRST + i
        vals = [i + 1, c["document"], c["norm_id"], c["actor"], c["unit"], c["text"],
                c["match"], c.get("actor_id") or "", c.get("label") or "",
                c.get("fallback_class") or "", None, None, None]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            if j >= 11:
                cell.fill = INPUT_FILL
    last = FIRST + len(cases) - 1
    dv_k = DataValidation(type="list", formula1='"yes,no,unsure"', allow_blank=True)
    dv_l = DataValidation(type="list", formula1=f"=Inventory!$A$2:$A${last_inv_row}",
                          allow_blank=True)
    dv_k.error, dv_l.error = "Choose yes, no or unsure.", "Choose an ID from the Inventory sheet."
    ws.add_data_validation(dv_k)
    ws.add_data_validation(dv_l)
    dv_k.add(f"K{FIRST}:K{last}")
    dv_l.add(f"L{FIRST}:L{last}")
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:M{last}"
    return last


def write_scores(wb, last: int, shares: dict[str, float], documents: list[str]) -> None:
    ws = wb.create_sheet("Scores")
    R = f"Review!$G${FIRST}:$G${last}"   # system match
    K = f"Review!$K${FIRST}:$K${last}"   # verdict
    L = f"Review!$L${FIRST}:$L${last}"   # correct answer
    B = f"Review!$B${FIRST}:$B${last}"   # document
    bold = Font(name=FONT, bold=True)
    norm = Font(name=FONT)
    blue = Font(name=FONT, color="0000FF")

    ws["A1"] = "Validation scores (update as the Review sheet is filled in)"
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    heads = ["stratum (system match)", "cases", "reviewed (yes+no)", "correct (yes)",
             "accuracy", "unsure", "share of all unit uses", "accuracy x share"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=3, column=j, value=h)
    _style_header(ws, 3, len(heads))
    for i, s in enumerate(STRATA):
        r = 4 + i
        ws.cell(row=r, column=1, value=s)
        ws.cell(row=r, column=2, value=f'=COUNTIF({R},A{r})')
        ws.cell(row=r, column=3, value=f'=COUNTIFS({R},A{r},{K},"yes")+COUNTIFS({R},A{r},{K},"no")')
        ws.cell(row=r, column=4, value=f'=COUNTIFS({R},A{r},{K},"yes")')
        ws.cell(row=r, column=5, value=f'=IF(C{r}=0,"",D{r}/C{r})')
        ws.cell(row=r, column=6, value=f'=COUNTIFS({R},A{r},{K},"unsure")')
        c = ws.cell(row=r, column=7, value=round(shares.get(s, 0.0), 4))
        c.font = blue
        c.comment = Comment("Share of this match type among all actor unit uses in "
                            "norm_actor_map.jsonl, computed by actor_validate.py when the "
                            "sample was drawn. Used to weight the overall score.", "actor_validate")
        ws.cell(row=r, column=8, value=f'=IF(C{r}=0,"",E{r}*G{r})')
    t = 8
    ws.cell(row=t, column=1, value="all (unweighted)")
    for col in "BCDF":
        ws[f"{col}{t}"] = f"=SUM({col}4:{col}7)"
    ws[f"E{t}"] = f'=IF(C{t}=0,"",D{t}/C{t})'
    ws[f"G{t}"] = "=SUM(G4:G7)"
    ws.cell(row=9, column=1, value="all (weighted to real shares)")
    ws["E9"] = '=IF(SUMIFS(G4:G7,C4:C7,">0")=0,"",SUMPRODUCT(E4:E7,G4:G7)/SUMIFS(G4:G7,C4:C7,">0"))'
    ws.cell(row=10, column=1, value="95% interval, unweighted (Wilson)")
    z = "1.96"
    p, n = "E8", "C8"
    ws["E10"] = (f'=IF({n}=0,"",({p}+{z}^2/(2*{n})-{z}*SQRT({p}*(1-{p})/{n}+{z}^2/(4*{n}^2)))'
                 f'/(1+{z}^2/{n}))')
    ws["F10"] = (f'=IF({n}=0,"",({p}+{z}^2/(2*{n})+{z}*SQRT({p}*(1-{p})/{n}+{z}^2/(4*{n}^2)))'
                 f'/(1+{z}^2/{n}))')
    ws["G10"] = "low / high"
    for r in range(4, 11):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            cell.font = blue if (c == 7 and r <= 7) else (bold if r >= 8 and c == 1 else norm)
            if c in (5, 8) or (r == 10 and c == 6):
                cell.number_format = "0.0%"
            if c == 7 and r <= 8:
                cell.number_format = "0.0%"

    # error kinds
    ws["A12"] = "Errors by kind"
    ws["A12"].font = bold
    rows = [
        ("wrong actor ID (exact/shortened, answer is another ID)",
         f'=COUNTIFS({R},"exact",{K},"no",{L},"<>fallback",{L},"<>not_actor",{L},"<>")'
         f'+COUNTIFS({R},"shortened",{K},"no",{L},"<>fallback",{L},"<>not_actor",{L},"<>")'),
        ("ID given, but should be fallback",
         f'=COUNTIFS({R},"exact",{K},"no",{L},"fallback")+COUNTIFS({R},"shortened",{K},"no",{L},"fallback")'),
        ("taken for an actor, but is not one",
         f'=COUNTIFS({R},"<>not_actor",{K},"no",{L},"not_actor")'),
        ("fallback, but an entry fits",
         f'=COUNTIFS({R},"fallback",{K},"no",{L},"<>not_actor",{L},"<>fallback",{L},"<>")'),
        ("not_actor, but is an actor (missed)",
         f'=COUNTIFS({R},"not_actor",{K},"no")'),
        ("'no' without a correct answer in column L",
         f'=COUNTIFS({K},"no",{L},"")'),
    ]
    for i, (label, f) in enumerate(rows):
        ws.cell(row=13 + i, column=1, value=label).font = norm
        ws.cell(row=13 + i, column=2, value=f).font = norm

    # per document
    r0 = 13 + len(rows) + 1
    ws.cell(row=r0, column=1, value="By document").font = bold
    for j, h in enumerate(["document", "cases", "reviewed (yes+no)", "correct (yes)", "accuracy"],
                          start=1):
        ws.cell(row=r0 + 1, column=j, value=h)
    _style_header(ws, r0 + 1, 5)
    for i, d in enumerate(documents):
        r = r0 + 2 + i
        ws.cell(row=r, column=1, value=d).font = norm
        ws.cell(row=r, column=2, value=f'=COUNTIF({B},A{r})').font = norm
        ws.cell(row=r, column=3, value=f'=COUNTIFS({B},A{r},{K},"yes")+COUNTIFS({B},A{r},{K},"no")').font = norm
        ws.cell(row=r, column=4, value=f'=COUNTIFS({B},A{r},{K},"yes")').font = norm
        c = ws.cell(row=r, column=5, value=f'=IF(C{r}=0,"",D{r}/C{r})')
        c.font, c.number_format = norm, "0.0%"
    ws.column_dimensions["A"].width = 52
    for letter in "BCDEFGH":
        ws.column_dimensions[letter].width = 16


def cmd_sample(args) -> None:
    d = args.dir
    units = [json.loads(l) for l in open(d / "unit_actor_map.jsonl", encoding="utf-8") if l.strip()]
    norms = [json.loads(l) for l in open(d / "norm_actor_map.jsonl", encoding="utf-8") if l.strip()]
    entries = [json.loads(l) for l in open(d / "actor_inventory.jsonl", encoding="utf-8") if l.strip()]
    texts = load_texts(args.records, args.id_field, args.text_field)

    # real shares of the strata, by norm uses
    uses = Counter()
    for u in units:
        uses[u["match"]] += u.get("norms", 1)
    total = sum(uses.values()) or 1
    shares = {s: uses[s] / total for s in STRATA}

    # norms that use each (document, unit), to show one of them
    users = defaultdict(list)
    for n in norms:
        for a in n.get("actors", []):
            users[(n["document"], a["unit"])].append(n)
        for u in n.get("dropped", []):
            users[(n["document"], u)].append(n)

    rng = random.Random(args.seed)
    q = quotas(args.size)
    picked = draw(units, q, rng)
    cases = []
    for u in picked:
        n = rng.choice(users[(u["document"], u["unit"])]) if users[(u["document"], u["unit"])] else {}
        text = texts.get(str(n.get("norm_id")), "")
        cases.append({**u, "norm_id": n.get("norm_id", ""), "actor": n.get("actor", ""),
                      "text": text[:1500]})
    rng.shuffle(cases)  # the reviewer should not see the strata in blocks

    wb = Workbook()
    write_instructions(wb, len(cases))
    last_inv = write_inventory(wb, entries)
    last = write_review(wb, cases, last_inv)
    write_scores(wb, last, shares, sorted({c["document"] for c in cases}))
    wb._sheets = [wb[n] for n in ("Instructions", "Review", "Scores", "Inventory")]
    wb.active = 1
    out = args.out or d / "actor_validation.xlsx"
    wb.save(out)
    got = Counter(c["match"] for c in cases)
    print(f"{len(cases)} cases -> {out}")
    print("per stratum: " + ", ".join(f"{s} {got[s]}/{q[s]}" for s in STRATA))
    print("per document: " + ", ".join(f"{k} {v}" for k, v in Counter(c['document'] for c in cases).items()))
    if not texts:
        print("No --records given: the norm text column is empty.")


# --------------------------------------------------------------------------
# Scoring
# --------------------------------------------------------------------------

def wilson(k: int, n: int, z: float = 1.96) -> tuple[float, float]:
    if n == 0:
        return float("nan"), float("nan")
    p = k / n
    centre = p + z * z / (2 * n)
    spread = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return (centre - spread) / (1 + z * z / n), (centre + spread) / (1 + z * z / n)


def cmd_score(args) -> None:
    wb = load_workbook(args.file, data_only=False)
    ws = wb["Review"]
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    rows = []
    for r in ws.iter_rows(min_row=FIRST, values_only=True):
        if r[col["case"]] is None:
            continue
        rows.append({"document": r[col["document"]], "match": r[col["system: match"]],
                     "verdict": (r[col["REVIEW: system correct?"]] or "").strip().lower(),
                     "answer": (r[col["REVIEW: correct answer (if no)"]] or "").strip(),
                     "unit": r[col["actor unit (cleaned)"]]})
    sc = wb["Scores"]
    shares = {sc.cell(row=4 + i, column=1).value: float(sc.cell(row=4 + i, column=7).value or 0)
              for i in range(len(STRATA))}

    def acc(sub):
        k = sum(1 for x in sub if x["verdict"] == "yes")
        n = sum(1 for x in sub if x["verdict"] in ("yes", "no"))
        return k, n

    report = {"cases": len(rows),
              "reviewed": sum(1 for x in rows if x["verdict"] in ("yes", "no")),
              "unsure": sum(1 for x in rows if x["verdict"] == "unsure"),
              "open": sum(1 for x in rows if x["verdict"] not in VERDICTS)}
    print(f"Cases {report['cases']}: reviewed {report['reviewed']}, unsure {report['unsure']}, "
          f"not filled in {report['open']}")
    print(f"\n{'stratum':<12}{'correct':>9}{'acc':>8}   95% interval   share")
    strata, wsum, wacc = {}, 0.0, 0.0
    for s in STRATA:
        k, n = acc([x for x in rows if x["match"] == s])
        lo, hi = wilson(k, n)
        strata[s] = {"correct": k, "reviewed": n, "accuracy": k / n if n else None,
                     "ci95": [lo, hi], "share": shares.get(s)}
        if n:
            wsum += shares.get(s, 0)
            wacc += shares.get(s, 0) * k / n
        print(f"{s:<12}{k:>4}/{n:<4}{(k / n if n else float('nan')):>8.1%}   "
              f"{lo:>5.1%} - {hi:<5.1%}  {shares.get(s, 0):>6.1%}")
    k, n = acc(rows)
    lo, hi = wilson(k, n)
    report["strata"] = strata
    report["accuracy_unweighted"] = {"correct": k, "reviewed": n,
                                     "accuracy": k / n if n else None, "ci95": [lo, hi]}
    report["accuracy_weighted"] = wacc / wsum if wsum else None
    print(f"{'all':<12}{k:>4}/{n:<4}{(k / n if n else float('nan')):>8.1%}   {lo:>5.1%} - {hi:<5.1%}")
    if wsum:
        print(f"weighted to real shares: {wacc / wsum:.1%}")

    mapped = [x for x in rows if x["match"] in ("exact", "shortened")]
    k, n = acc(mapped)
    report["id_precision"] = k / n if n else None
    is_actor = [x for x in rows if x["match"] != "not_actor" and x["verdict"] in ("yes", "no")]
    false_actor = sum(1 for x in is_actor if x["verdict"] == "no" and x["answer"] == "not_actor")
    report["actor_detection_precision"] = 1 - false_actor / len(is_actor) if is_actor else None
    print(f"\nactor ID precision (exact+shortened): "
          f"{report['id_precision']:.1%}" if n else "\nactor ID precision: no reviewed cases")
    if is_actor:
        print(f"units taken for actors that are actors: {report['actor_detection_precision']:.1%}")

    kinds = Counter()
    for x in rows:
        if x["verdict"] != "no":
            continue
        a, m = x["answer"], x["match"]
        if not a:
            kinds["'no' without a correct answer"] += 1
        elif m == "not_actor":
            kinds["missed actor (not_actor, but is an actor)"] += 1
        elif a == "not_actor":
            kinds["taken for an actor, but is not one"] += 1
        elif a == "fallback":
            kinds["ID given, but should be fallback"] += 1
        elif m == "fallback":
            kinds["fallback, but an entry fits"] += 1
        else:
            kinds["wrong actor ID"] += 1
    report["errors"] = dict(kinds)
    if kinds:
        print("\nerrors by kind:")
        for kname, v in kinds.most_common():
            print(f"  {v:>3}  {kname}")
    by_doc = {}
    for dname in sorted({x["document"] for x in rows}):
        k, n = acc([x for x in rows if x["document"] == dname])
        by_doc[dname] = {"correct": k, "reviewed": n, "accuracy": k / n if n else None}
    report["by_document"] = by_doc
    print("\nby document: " + ", ".join(
        f"{dname} {v['correct']}/{v['reviewed']}" for dname, v in by_doc.items()))

    out = Path(str(args.file) + ".scores.json")
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n-> {out}")


# --------------------------------------------------------------------------

def main() -> None:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample", help="draw the sample and write the review workbook")
    s.add_argument("--dir", type=Path, default=Path("data/conflict/actor"))
    s.add_argument("--records", type=Path, default=None, help="norm records with the norm text")
    s.add_argument("--id-field", default="norm_id")
    s.add_argument("--text-field", default="norm.text")
    s.add_argument("--size", type=int, default=40, help="number of cases, 20-50")
    s.add_argument("--seed", type=int, default=13)
    s.add_argument("--out", type=Path, default=None)
    c = sub.add_parser("score", help="score a filled review workbook")
    c.add_argument("--file", type=Path, default=Path("data/conflict/actor/actor_validation.xlsx"))
    args = p.parse_args()
    cmd_sample(args) if args.cmd == "sample" else cmd_score(args)


if __name__ == "__main__":
    main()