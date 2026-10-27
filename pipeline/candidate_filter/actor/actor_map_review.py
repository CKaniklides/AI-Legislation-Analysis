#!/usr/bin/env python3
"""actor_map_review.py - validation of the actor mapping (actor filter, Step 2b).

Two steps:
  1. sample  - draw a stratified sample of actor units as they occur in norms and write a
               workbook for blind labelling: is the actor the unit was mapped to right?
  2. review  - read the labels back and report precision per stratum, per source and per
               match level, the EU - Dutch difference, actors lost as 'not an actor',
               missed inventory entries and the kinds of error, with 95% intervals.

What is validated
  actor_map.py gives every actor unit of a norm one of four results, each a claim:
    exact      the unit is this inventory actor (its full role key is the entry's key)
    shortened  the unit is this inventory actor (found after dropping words)
    fallback   the unit is an actor of this broad class, with no fitting inventory entry
    not_actor  the unit is not an actor (it is dropped from the norm)

Population and sampling unit
  An actor unit as it occurs in a norm: the kept actors and the dropped units of
  norm_actor_map.jsonl. A unit used in 90 norms counts 90 times, as it does downstream.
  Norms with status pronoun, missing or no_units have no unit and make no mapping claim;
  they are exact counts, by source.

Census of the most frequent results
  The result of a unit text is fixed wherever it gets the same actor ID, class or
  'not an actor', so one judgement settles all those occurrences, and an error repeats in
  every norm that uses them. The --census (20) most frequent (unit, result) groups are
  judged in full, each once, against up to three contexts taken in turn from both
  sources. They are the top of the frequency ranking (ties alphabetical), fixed before
  any judging. The coverage curve is in the Design sheet.

Strata outside the census: source (EU, NL) x match level (exact, shortened, fallback,
  not_actor). Each source gets --per-source (50) judgements: a floor of --floor (8) per
  stratum and the rest in proportion to N_h. A norm gives at most --max-per-norm (1)
  sampled unit. Default total: 20 + 50 + 50 = 120.

Consistency
  If actor_map.py is found (next to this script, or --mapper), every distinct unit is
  mapped again with the current inventory. The script stops if more than --max-mismatch
  (1%) of the results differ: the map is stale or the inventory has changed.

Labels: correct / partial / wrong / unclear; for partial or wrong, 'What is wrong?' and
'Correct actor'. Strict precision counts correct, lenient counts correct and partial.

Inputs (in --data)
  norm_actor_map.jsonl, unit_actor_map.jsonl, actor_inventory.jsonl (actor_map.py and
  actor_inventory.py); --norm-text for the norm texts (JSON array or JSONL).
Outputs
  sample: actor_map_review.xlsx (or --out)
  review: printed, and <workbook>.report.json

Usage
  python actor_map_review.py sample --data data/conflict/actor \\
      --norm-text data/conflict/norms_all.json
  python actor_map_review.py review --xlsx data/conflict/actor/actor_map_review.xlsx
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import math
import random
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SOURCES = ["EU", "NL"]
LEVELS = ["exact", "shortened", "fallback", "not_actor"]
CLAIM = {"exact": "actor ID", "shortened": "actor ID", "fallback": "broad class",
         "not_actor": "not an actor"}
CLAIM_LEVELS = {"actor ID": ["exact", "shortened"], "broad class": ["fallback"],
                "not an actor": ["not_actor"]}
LABELS = ["correct", "partial", "wrong", "unclear"]
PROBLEMS = ["other actor", "too broad", "too narrow", "missed entry", "wrong class",
            "not an actor", "actor, not in inventory", "badly cut unit"]
NO_UNIT = ["pronoun", "missing", "no_units"]
CENSUS = "census"
STRATA = [f"{s} | {l}" for s in SOURCES for l in LEVELS]

WHY = {
    "exact": "The full role key equals an inventory label or alias: the safest route. Checks "
             "the inventory entries of the common actors.",
    "shortened": "Words were dropped until a key matched: the riskiest route to an ID (an "
                 "entry too broad, or the wrong one).",
    "fallback": "No entry matched; only the head word's class. Is it an actor, is the class "
                "right, and was an inventory entry missed?",
    "not_actor": "Dropped as not an actor. Its errors are actors lost to the whole actor "
                 "filter (recall of the mapping).",
}

Z = 1.96
FONT = "Arial"
INPUT_FILL = PatternFill("solid", start_color="FFFF00")
HEAD_FILL = PatternFill("solid", start_color="1F3864")
THIN = Side(style="thin", color="BFBFBF")


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------

def _get(obj, dotted):
    for k in dotted.split("."):
        if not isinstance(obj, dict):
            return None
        obj = obj.get(k)
    return obj


def read_jsonl(path):
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def read_records(path):
    raw = Path(path).read_text(encoding="utf-8").strip()
    return json.loads(raw) if raw.startswith("[") else \
        [json.loads(l) for l in raw.splitlines() if l.strip()]


def source_of(document):
    return "NL" if str(document).upper().startswith("BWBR") else "EU"


def result_of(rec):
    """What the mapping claims: an actor ID, a broad class, or 'not an actor'."""
    if rec.get("actor_id"):
        return rec["actor_id"]
    if rec.get("fallback_class"):
        return f"class:{rec['fallback_class']}"
    return "not_actor"


def level_of(rec):
    m = rec.get("match")
    if m in LEVELS:
        return m
    return "fallback" if rec.get("fallback_class") else "not_actor"


def population(norms, unit_map):
    """Unit occurrences (kept actors and dropped units) and exact counts of the rest."""
    occ, norm_status = [], Counter()
    for n in norms:
        doc = str(n.get("document") or str(n.get("norm_id", "")).split(":", 1)[0])
        src = source_of(doc)
        norm_status[(src, n.get("status") or "?")] += 1
        recs = [dict(a) for a in n.get("actors") or []]
        for u in n.get("dropped") or []:
            r = dict(unit_map.get((doc, u)) or {"unit": u})
            r.update(match="not_actor", actor_id=None, fallback_class=None)
            recs.append(r)
        for k, r in enumerate(recs):
            lvl = level_of(r)
            occ.append({"norm_id": str(n.get("norm_id")), "document": doc, "source": src,
                        "actor": n.get("actor") or "", "pos": k, "unit": r.get("unit") or "",
                        "level": lvl, "result": result_of(r), "actor_id": r.get("actor_id"),
                        "label": r.get("label"), "fallback_class": r.get("fallback_class"),
                        "role_keys": " / ".join(r.get("role_keys") or [])})
    return occ, norm_status


def load_mapper(args):
    """actor_map.py as a module, or None (then the recheck is skipped)."""
    path = Path(args.mapper) if args.mapper else Path(__file__).with_name("actor_map.py")
    if not path.is_file():
        print(f"Note: {path} not found; the mapping is not rechecked.")
        return None
    sys.path.insert(0, str(path.parent))   # actor_map imports actor_lexicon from its folder
    try:
        spec = importlib.util.spec_from_file_location("actor_map", path)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
    except Exception as e:  # missing nltk, actor_heads.tsv, ...
        print(f"Note: {path} could not be loaded ({e}); the mapping is not rechecked.")
        return None
    return mod


def recheck(mod, entries, occ, max_mismatch):
    """Map every distinct (document, unit) again; stop if too many results differ."""
    index = mod.index_inventory(entries)
    seen, bad = {}, []
    for o in occ:
        k = (o["document"], o["unit"])
        if k in seen:
            continue
        new = mod.map_unit(o["unit"], o["document"], index, Counter())
        seen[k] = (result_of(new), level_of(new)) == (o["result"], o["level"])
        if not seen[k]:
            bad.append((k, o["result"], result_of(new)))
    share = len(bad) / max(1, len(seen))
    if bad:
        msg = (f"The current mapping differs from norm_actor_map.jsonl for {len(bad):,} of "
               f"{len(seen):,} distinct units ({share:.1%}), e.g. "
               + "; ".join(f"{u!r} in {d}: {a} -> {b}" for (d, u), a, b in bad[:3]) + ".")
        if share > max_mismatch:
            raise SystemExit(msg + f" That is more than --max-mismatch ({max_mismatch:.0%}): "
                             "re-run actor_map.py with the current inventory.")
        print("Warning: " + msg)
    return len(bad), len(seen)


# ---------------------------------------------------------------------------
# Statistics
# ---------------------------------------------------------------------------

def wilson(p, n, z=Z):
    if n <= 0:
        return 0.0, 1.0
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(max(0.0, p * (1 - p) / n + z * z / (4 * n * n))) / d
    return max(0.0, c - h), min(1.0, c + h)


def estimate(parts, exact=(0, 0.0, 0)):
    """parts: sampled strata (N_h, k, n); exact: census (occurrences, agreeing occurrences,
    items). Stratified proportion; Wilson on the effective n, capped at the labelled items."""
    parts = [x for x in parts if x[0] > 0 and x[2] > 0]
    n_ex, k_ex, items = exact
    total = sum(w for w, _, _ in parts) + n_ex
    if total <= 0:
        return None
    p = (sum(w * k / n for w, k, n in parts) + k_ex) / total
    var = sum((w / total) ** 2 * max(0.0, 1 - n / w) * (k / n) * (1 - k / n) / n
              for w, k, n in parts)
    n_all = sum(n for _, _, n in parts) + items
    n_eff = min(n_all, p * (1 - p) / var) if var > 0 else n_all
    lo, hi = wilson(p, n_eff)
    return {"p": p, "lo": lo, "hi": hi, "n_eff": n_eff, "var": var, "n": n_all,
            "weight": total}


def newcombe(a, b):
    """Difference a - b with Newcombe's hybrid interval from the two intervals."""
    d = a["p"] - b["p"]
    lo = d - math.sqrt((a["p"] - a["lo"]) ** 2 + (b["hi"] - b["p"]) ** 2)
    hi = d + math.sqrt((a["hi"] - a["p"]) ** 2 + (b["p"] - b["lo"]) ** 2)
    return {"d": d, "lo": lo, "hi": hi}


def fmt(res):
    if res is None:
        return "(no labels)"
    return f"{res['p']:5.1%} [{res['lo']:5.1%} - {res['hi']:5.1%}]   n_eff {res['n_eff']:5.1f}"


def allocate(N, budget, floor):
    """A floor per stratum, the rest in proportion to N_h (D'Hondt), never above N_h."""
    alloc = {s: min(floor, N[s]) for s in N}
    for _ in range(max(0, budget - sum(alloc.values()))):
        open_ = [s for s in N if alloc[s] < N[s]]
        if not open_:
            break
        best = max(open_, key=lambda s: (N[s] / (alloc[s] - min(floor, N[s]) + 1), s))
        alloc[best] += 1
    return alloc


def expected_halfwidth(N_src, alloc, census_occ, p):
    tot = sum(N_src.values()) + census_occ
    if not tot:
        return float("nan")
    v = sum((N_src[s] / tot) ** 2 * max(0.0, 1 - alloc[s] / N_src[s]) * p * (1 - p) / alloc[s]
            for s in N_src if alloc[s] and N_src[s])
    return 100 * Z * math.sqrt(v)


# ---------------------------------------------------------------------------
# Sampling
# ---------------------------------------------------------------------------

def sample(args):
    data = Path(args.data)
    out = Path(args.out) if args.out else data / "actor_map_review.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)

    paths = {}
    for name, what in ((args.norms, "--norms"), (args.units, "--units"),
                       (args.inventory, "--inventory")):
        p = Path(name) if Path(name).is_file() else data / name
        if not p.is_file():
            raise SystemExit(f"{what}: {p} not found.")
        paths[what] = p
    norms = read_jsonl(paths["--norms"])
    unit_map = {(str(r["document"]), r["unit"]): r for r in read_jsonl(paths["--units"])}
    entries = read_jsonl(paths["--inventory"])
    by_id = {e["id"]: e for e in entries}

    occ, norm_status = population(norms, unit_map)
    if not occ:
        raise SystemExit("No actor units in the norm map.")
    unknown_ids = sorted({o["actor_id"] for o in occ if o["actor_id"] and o["actor_id"] not in by_id})
    if unknown_ids:
        raise SystemExit(f"{len(unknown_ids)} actor IDs of the map are not in the inventory, e.g. "
                         f"{unknown_ids[:3]}: the map was made with another inventory. "
                         "Re-run actor_map.py.")

    mod = load_mapper(args)
    mismatch = recheck(mod, entries, occ, args.max_mismatch) if mod else None

    # --- census: the most frequent (unit, result) groups
    groups = defaultdict(list)
    for o in occ:
        groups[(o["unit"], o["result"])].append(o)
    ranking = sorted(groups, key=lambda g: (-len(groups[g]), g[0].lower(), g[1]))
    census = ranking[:args.census]
    in_census = set(census)
    src_total = Counter(o["source"] for o in occ)
    coverage = {}
    for k in (5, 10, 20, 30, 50, 100):
        c = Counter(o["source"] for g in ranking[:k] for o in groups[g])
        coverage[k] = {s: c[s] / src_total[s] if src_total[s] else 0.0 for s in SOURCES}

    # --- strata outside the census
    pools = defaultdict(list)
    for o in occ:
        if (o["unit"], o["result"]) not in in_census:
            pools[f"{o['source']} | {o['level']}"].append(o)
    N = {s: len(pools[s]) for s in STRATA}
    alloc = {}
    for src in SOURCES:
        names = [f"{src} | {l}" for l in LEVELS]
        alloc.update(allocate({s: N[s] for s in names}, args.per_source, args.floor))
    for item in args.alloc or []:
        name, _, n = item.rpartition("=")
        if name.strip() not in alloc:
            raise SystemExit(f"--alloc: unknown stratum '{name.strip()}'. Strata: {STRATA}")
        alloc[name.strip()] = min(int(n), N[name.strip()])

    rng = random.Random(args.seed)
    uses = Counter()
    rows = []
    # census rows: up to three contexts, taken in turn from the sources
    for g in census:
        by_src = {s: sorted((o for o in groups[g] if o["source"] == s),
                            key=lambda o: (o["norm_id"], o["pos"])) for s in SOURCES}
        for s in SOURCES:
            rng.shuffle(by_src[s])
        ctx, seen = [], set()
        while len(ctx) < args.contexts and any(by_src.values()):
            for s in SOURCES:
                while by_src[s] and len(ctx) < args.contexts:
                    o = by_src[s].pop(0)
                    if o["norm_id"] not in seen:
                        seen.add(o["norm_id"])
                        ctx.append(o)
                        break
        for o in ctx:
            uses[o["norm_id"]] += 1
        first = groups[g][0]
        cnt = Counter(o["source"] for o in groups[g])
        rows.append({"stratum": CENSUS, "kind": "census",
                     "source": "+".join(s for s in SOURCES if cnt[s]),
                     "level": first["level"], "unit": g[0], "result": g[1],
                     "actor_id": first["actor_id"], "fallback_class": first["fallback_class"],
                     "role_keys": first["role_keys"], "contexts": ctx,
                     "occ_EU": cnt["EU"], "occ_NL": cnt["NL"], "group_occ": len(groups[g])})
    # sampled rows: smallest strata first, at most --max-per-norm units per norm
    for s in sorted(STRATA, key=lambda s: N[s]):
        pool = sorted(pools[s], key=lambda o: (o["norm_id"], o["pos"]))
        rng.shuffle(pool)
        taken = 0
        for o in pool:
            if taken == alloc[s]:
                break
            if uses[o["norm_id"]] >= args.max_per_norm:
                continue
            uses[o["norm_id"]] += 1
            taken += 1
            g = (o["unit"], o["result"])
            rows.append({"stratum": s, "kind": "sample", "source": o["source"],
                         "level": o["level"], "unit": o["unit"], "result": o["result"],
                         "actor_id": o["actor_id"], "fallback_class": o["fallback_class"],
                         "role_keys": o["role_keys"], "contexts": [o],
                         "occ_EU": 0, "occ_NL": 0, "group_occ": len(groups[g])})
        if taken < alloc[s]:
            print(f"Warning: '{s}' yielded {taken} units (target {alloc[s]}).")
    rng.shuffle(rows)   # order must not reveal the stratum or the census

    texts = {}
    if args.norm_text:
        for r in read_records(args.norm_text):
            texts[str(_get(r, args.id_field))] = str(_get(r, args.text_field) or "")

    info = dict(N=N, alloc=alloc, taken=Counter(r["stratum"] for r in rows if r["kind"] == "sample"),
                census=[(g, Counter(o["source"] for o in groups[g]), groups[g][0]["level"])
                        for g in census],
                coverage=coverage, src_total=src_total, norm_status=norm_status,
                level_total=Counter((o["source"], o["level"]) for o in occ),
                mismatch=mismatch, n_norms=len(norms), args=args)
    write_workbook(out, rows, by_id, entries, texts, info)

    print(f"Wrote {len(rows)} rows ({len(census)} census, {len(rows) - len(census)} sampled) -> {out}")
    print(f"Unit occurrences {len(occ):,}: EU {src_total['EU']:,}, NL {src_total['NL']:,}; "
          f"census covers {sum(len(groups[g]) for g in census):,}")
    print(f"{'stratum':22} {'N_h':>8} {'n':>4}")
    for s in STRATA:
        print(f"{s:22} {N[s]:>8,} {info['taken'][s]:>4}")
    if mismatch and mismatch[0]:
        print(f"Recheck: {mismatch[0]:,} of {mismatch[1]:,} distinct units differ.")
    elif mismatch is not None:
        print(f"Recheck: all {mismatch[1]:,} distinct units map as in the map.")
    if not texts:
        print("No --norm-text given: the text column is empty.")


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def header(ws, row, ncols, col0=1):
    for c in range(col0, col0 + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def put(ws, row, col, value, wrap=False, bold=False, italic=False, numfmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, bold=bold, italic=italic)
    c.alignment = Alignment(vertical="top", wrap_text=wrap)
    if numfmt:
        c.number_format = numfmt
    return c


def assigned(row, by_id):
    """What the labeller sees of the result: no match level, no role keys."""
    if row["actor_id"]:
        e = by_id.get(row["actor_id"], {})
        return f"{e.get('label', row['actor_id'])}   [{row['actor_id']}]"
    if row["fallback_class"]:
        return f"an actor of class '{row['fallback_class']}' (no inventory entry)"
    return "not an actor (dropped)"


def describe(e):
    if not e:
        return ""
    if e.get("kind") == "defined":
        return f"Defined in {e.get('source')}: {e.get('definition') or ''}"[:700]
    extra = f" Aliases: {'; '.join(e['aliases'])}." if e.get("aliases") else ""
    part = f" Part of {e['part_of']}." if e.get("part_of") else ""
    scope = e.get("scope")
    scope = scope if isinstance(scope, str) else ", ".join(scope or [])
    return f"Known actor (not defined), class {e.get('class')}, valid in {scope}.{extra}{part}"


def design_note(info):
    a, N = info["args"], info["N"]
    census_occ = {s: sum(c[s] for _, c, _ in info["census"]) for s in SOURCES}
    hw = {}
    for s in SOURCES:
        names = {x: N[x] for x in STRATA if x.startswith(s)}
        hw[s] = (expected_halfwidth(names, info["alloc"], census_occ[s], 0.5),
                 expected_halfwidth(names, info["alloc"], census_occ[s], 0.8))
    diff50 = math.sqrt(hw["EU"][0] ** 2 + hw["NL"][0] ** 2)
    mm = info["mismatch"]
    recheck_txt = ("actor_map.py was not available, so the map was not rechecked against the "
                   "current inventory." if mm is None else
                   f"Every distinct unit was mapped again with the current inventory; "
                   f"{mm[0]:,} of {mm[1]:,} differed." if mm[0] else
                   f"Every distinct unit ({mm[1]:,}) was mapped again with the current inventory "
                   "and gave the same result.")
    return (
        "What is validated. Each actor unit of a norm is mapped to an inventory actor (exact or "
        "shortened match), to the broad class of its head word (fallback) or to nothing (not an "
        "actor, dropped). Each result is a claim the labeller checks from the norm text: is this "
        "the actor, is it an actor of this class without an entry, is it really not an actor?\n\n"
        "Population. The sampling unit is an actor unit as it occurs in a norm, so a unit used "
        "in many norms counts as often as it is used downstream, and each is judged in its own "
        f"norm. {sum(info['src_total'].values()):,} unit occurrences: EU "
        f"{info['src_total']['EU']:,}, Dutch {info['src_total']['NL']:,}. Norms with status "
        "pronoun, missing or no_units have no unit and make no mapping claim; they are exact "
        "counts (table below). The source is taken from the document id: BWBR ids are Dutch, "
        "all others EU.\n\n"
        f"Census. The result of a unit is the same wherever it gets the same ID, class or 'not "
        f"an actor', so one judgement settles all those occurrences. The {a.census} most "
        "frequent (unit, result) groups are judged in full, each once, against up to "
        f"{a.contexts} contexts taken in turn from both sources. They are the top of the "
        "frequency ranking, ties broken alphabetically, fixed before any judging. A census "
        "group that fits some uses but not others is labelled partial. The coverage curve "
        "below shows where a further group would cover little more than one sampled unit.\n\n"
        "Strata. Outside the census, units are stratified by source, so the EU-Dutch "
        "comparison gets the same number of judgements on each side, and by match level, "
        "because each level makes a different claim and can fail at its own rate. Each source "
        f"gets {a.per_source} judgements: a floor of {a.floor} per stratum, the rest in "
        f"proportion to N_h. A norm gives at most {a.max_per_norm} sampled unit.\n\n"
        f"Precision. In the worst case (50% in every stratum) each source's precision is known "
        f"to about +/-{hw['EU'][0]:.0f} points (EU) and +/-{hw['NL'][0]:.0f} (Dutch), at 80% to "
        f"+/-{hw['EU'][1]:.0f} and +/-{hw['NL'][1]:.0f}; the EU-Dutch difference to about "
        f"+/-{diff50:.0f} points. The census adds no sampling error. Overall figures weight "
        "each stratum by N_h and each census group by its occurrences, never by row count.\n\n"
        "Blind. The labeller sees the norm, the unit and what it was mapped to, but not the "
        "match level or the role keys; those are in the hidden key sheet, and the rows are "
        "shuffled. A fallback or 'not an actor' result is visible as such, since it is the "
        "claim being judged.\n\n"
        f"Consistency. {recheck_txt} Adding actors to the inventory changes the strata, so "
        "the sample is drawn after the inventory is updated."
    )


def write_workbook(path, rows, by_id, entries, texts, info):
    a, N = info["args"], info["N"]
    wb = Workbook()

    # ---------------- Instructions ----------------
    g = wb.active
    g.title = "Instructions"
    lines = [
        ("Actor mapping review", True),
        ("Each row of the Review sheet is an actor unit in a norm (census rows show up to "
         f"{a.contexts} norms that use the same unit). 'Assigned' is what the system mapped the "
         "unit to. Judge whether that is right, from the norm text and the actor phrase. The "
         "Inventory sheet lists every known actor with its definition.", False),
        ("Assessment (yellow)", True),
        ("Assigned an actor [ID]:", True),
        ("  correct    the unit is this actor.", False),
        ("  partial    right actor family but too broad or too narrow (another entry fits "
         "better), or right for some census contexts only.", False),
        ("  wrong      another actor, or not an actor at all.", False),
        ("Assigned 'an actor of class X (no inventory entry)':", True),
        ("  correct    an actor, no inventory entry fits, and class X fits.", False),
        ("  partial    an actor and class X fits, but an inventory entry fits (missed entry).", False),
        ("  wrong      the class does not fit, or it is not an actor.", False),
        ("Assigned 'not an actor (dropped)':", True),
        ("  correct    indeed not an actor (a document, decision, system, abstraction ...).", False),
        ("  wrong      it is an actor (use 'partial' only for census groups that are actors "
         "in some contexts).", False),
        ("  unclear    cannot tell from the text (not counted).", False),
        ("What is wrong? (for partial or wrong)", True),
        ("  other actor / too broad / too narrow / missed entry (an inventory entry fits) / "
         "wrong class / not an actor / actor, not in inventory (a real actor that has no "
         "entry) / badly cut unit (the unit is a fragment or several actors in one).", False),
        ("Correct actor: the inventory label or ID that fits, or the name of the actor if it "
         "has no entry. This feeds the list of inventory additions.", False),
        ("The hidden 'key' sheet holds the system's details: leave it alone until you are done.", False),
        ("When finished, save and run: python actor_map_review.py review --xlsx <this file>", False),
        ("", False),
        ("Example (format only; not part of the sample)", True),
        ("Unit: de Autoriteit Consument en Markt | Assigned: not an actor (dropped) | "
         "Assessment: wrong | What is wrong?: actor, not in inventory | "
         "Correct actor: Autoriteit Consument en Markt", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        put(g, i, 1, text, bold=bold, wrap=True)
    g.column_dimensions["A"].width = 120

    # ---------------- Review (blind) ----------------
    ws = wb.create_sheet("Review")
    cols = [("#", 5), ("Norm(s)", 20), ("Actor phrase", 34), ("Unit", 28)]
    if texts:
        cols += [("Text", 60)]
    cols += [("Assigned", 34), ("Assigned: definition or source", 50),
             ("Assessment", 13), ("What is wrong?", 20), ("Correct actor", 26), ("Note", 26)]
    names = [c for c, _ in cols]
    for j, (name, width) in enumerate(cols, 1):
        ws.cell(row=1, column=j, value=name)
        ws.column_dimensions[get_column_letter(j)].width = width
    header(ws, 1, len(cols))
    per_ctx = max(300, a.text_chars // max(1, a.contexts))
    for r, row in enumerate(rows, 2):
        ctx = row["contexts"]
        multi = len(ctx) > 1
        def joined(f):
            return "\n".join((f"[{k + 1}] " if multi else "") + f(o) for k, o in enumerate(ctx))
        values = {"#": r - 1, "Norm(s)": joined(lambda o: o["norm_id"]),
                  "Actor phrase": joined(lambda o: o["actor"]), "Unit": row["unit"],
                  "Text": "\n\n".join((f"[{k + 1}] " if multi else "")
                                      + texts.get(o["norm_id"], "")[:per_ctx if multi else a.text_chars]
                                      for k, o in enumerate(ctx)),
                  "Assigned": assigned(row, by_id),
                  "Assigned: definition or source": describe(by_id.get(row["actor_id"]))}
        for j, name in enumerate(names, 1):
            c = put(ws, r, j, values.get(name), wrap=name != "#")
            c.border = Border(bottom=THIN)
            if name in ("Assessment", "What is wrong?", "Correct actor", "Note"):
                c.fill = INPUT_FILL
    last = len(rows) + 1
    acol = get_column_letter(names.index("Assessment") + 1)
    pcol = get_column_letter(names.index("What is wrong?") + 1)
    for col, options in ((acol, LABELS), (pcol, PROBLEMS)):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    ws.freeze_panes = "B2"

    # ---------------- Inventory (reference) ----------------
    inv = wb.create_sheet("Inventory")
    ih = ["id", "label", "aliases", "kind", "valid in", "class", "part of", "definition or source"]
    inv.append(ih)
    header(inv, 1, len(ih))
    for e in entries:
        scope = e.get("documents") if e.get("kind") == "defined" else e.get("scope")
        scope = scope if isinstance(scope, str) else ", ".join(scope or [])
        inv.append([e["id"], e.get("label"), "; ".join(e.get("aliases") or []), e.get("kind"),
                    scope, e.get("class"), e.get("part_of"),
                    (e.get("definition") or e.get("source") or "")[:500]])
    for row in inv.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for col, w in zip("ABCDEFGH", [34, 30, 30, 9, 16, 12, 24, 70]):
        inv.column_dimensions[col].width = w
    inv.auto_filter.ref = f"A1:H{len(entries) + 1}"
    inv.freeze_panes = "C2"

    # ---------------- key (hidden) ----------------
    ks = wb.create_sheet("key")
    kh = ["#", "stratum", "kind", "source", "level", "unit", "result", "actor_id",
          "fallback_class", "role_keys", "occ_EU", "occ_NL", "group_occ", "norm_ids"]
    ks.append(kh)
    for i, row in enumerate(rows, 1):
        ks.append([i, row["stratum"], row["kind"], row["source"], row["level"], row["unit"],
                   row["result"], row["actor_id"], row["fallback_class"], row["role_keys"],
                   row["occ_EU"], row["occ_NL"], row["group_occ"],
                   ", ".join(o["norm_id"] for o in row["contexts"])])
    for row in ks.iter_rows():
        for c in row:
            c.font = Font(name=FONT)
    ks.sheet_state = "hidden"

    # ---------------- weights (hidden; read by 'review') ----------------
    wsx = wb.create_sheet("weights")
    wsx.append(["stratum", "source", "level", "N_h"])
    for s in STRATA:
        src, lvl = s.split(" | ")
        wsx.append([s, src, lvl, N[s]])
    for row in wsx.iter_rows():
        for c in row:
            c.font = Font(name=FONT)
    wsx.sheet_state = "hidden"

    # ---------------- Design ----------------
    ds = wb.create_sheet("Design")
    put(ds, 1, 1, "Sample design", bold=True)
    ds.merge_cells("A2:G2")
    put(ds, 2, 1, design_note(info), wrap=True)
    ds.row_dimensions[2].height = 520
    th = ["stratum", "source", "level", "population N_h", "share of source", "n sampled",
          "why this stratum and this n"]
    hr = 4
    for j, h in enumerate(th, 1):
        ds.cell(row=hr, column=j, value=h)
    header(ds, hr, len(th))
    first = hr + 1
    for k, s in enumerate(STRATA):
        r = first + k
        src, lvl = s.split(" | ")
        put(ds, r, 1, s)
        put(ds, r, 2, src)
        put(ds, r, 3, lvl)
        put(ds, r, 4, N[s], numfmt="#,##0")
        put(ds, r, 5, f'=IFERROR(D{r}/SUMIFS($D${first}:$D${first + len(STRATA) - 1},'
                      f'$B${first}:$B${first + len(STRATA) - 1},B{r}),0)', numfmt="0.0%")
        put(ds, r, 6, info["taken"][s])
        put(ds, r, 7, WHY[lvl], wrap=True)
    lastrow = first + len(STRATA) - 1
    tr = lastrow + 1
    put(ds, tr, 1, "total outside the census", bold=True)
    put(ds, tr, 4, f"=SUM(D{first}:D{lastrow})", bold=True, numfmt="#,##0")
    put(ds, tr, 6, f"=SUM(F{first}:F{lastrow})", bold=True)

    r = tr + 2
    put(ds, r, 1, f"Census: the {len(info['census'])} most frequent (unit, result) groups, "
                  "judged once each", bold=True)
    ch = ["#", "unit", "assigned", "level", "EU", "NL"]
    for j, h in enumerate(ch, 1):
        ds.cell(row=r + 1, column=j, value=h)
    header(ds, r + 1, len(ch))
    for k, ((unit, res), cnt, lvl) in enumerate(info["census"], 1):
        rr = r + 1 + k
        for j, v in enumerate([k, unit, res, lvl, cnt["EU"], cnt["NL"]], 1):
            put(ds, rr, j, v, numfmt="#,##0" if j >= 5 else None)
    cend = r + 1 + len(info["census"])
    put(ds, cend + 1, 1, "total", bold=True)
    put(ds, cend + 1, 5, f"=SUM(E{r + 2}:E{cend})", bold=True, numfmt="#,##0")
    put(ds, cend + 1, 6, f"=SUM(F{r + 2}:F{cend})", bold=True, numfmt="#,##0")

    r = cend + 3
    put(ds, r, 1, "Coverage: share of unit occurrences covered by the most frequent groups", bold=True)
    for j, h in enumerate(["census size", "EU", "NL"], 1):
        ds.cell(row=r + 1, column=j, value=h)
    header(ds, r + 1, 3)
    for k, (size, cov) in enumerate(info["coverage"].items()):
        rr = r + 2 + k
        put(ds, rr, 1, size)
        put(ds, rr, 2, cov["EU"], numfmt="0.0%")
        put(ds, rr, 3, cov["NL"], numfmt="0.0%")

    r = r + 2 + len(info["coverage"]) + 1
    put(ds, r, 1, "Exact counts (not sampled)", bold=True)
    for j, h in enumerate(["", "EU", "NL", "total"], 1):
        ds.cell(row=r + 1, column=j, value=h)
    header(ds, r + 1, 4)
    ns, lt = info["norm_status"], info["level_total"]
    statuses = sorted({st for _, st in ns}, key=lambda s: (s in NO_UNIT, s))
    exact = [("norms", None)] + [(f"   ... status {st}" + (" (no unit, no claim)" if st in NO_UNIT else ""),
                                  ("status", st)) for st in statuses] + \
            [("unit occurrences", None)] + [(f"   ... {lvl}", ("level", lvl)) for lvl in LEVELS]
    rr = r + 2
    for label, keyx in exact:
        put(ds, rr, 1, label, bold=keyx is None)
        if keyx:
            vals = [(ns if keyx[0] == "status" else lt)[(s, keyx[1])] for s in SOURCES]
            put(ds, rr, 2, vals[0], numfmt="#,##0")
            put(ds, rr, 3, vals[1], numfmt="#,##0")
            put(ds, rr, 4, f"=B{rr}+C{rr}", numfmt="#,##0")
        rr += 1
    put(ds, rr, 1, f"Settings: census {a.census}, contexts {a.contexts}, per source "
                   f"{a.per_source}, floor {a.floor}, max per norm {a.max_per_norm}, seed {a.seed}",
        italic=True)
    for col, w in zip("ABCDEFG", [44, 34, 22, 14, 12, 10, 70]):
        ds.column_dimensions[col].width = w

    # ---------------- Summary (live, unweighted per stratum) ----------------
    sm = wb.create_sheet("Summary")
    sh = ["stratum", "rows", "judged", "correct", "partial", "wrong", "unclear",
          "strict", "lenient", "strict 95% low", "strict 95% high"]
    sm.append(sh)
    header(sm, 1, len(sh))
    sm.cell(row=1, column=13, value="z").font = Font(name=FONT, bold=True)
    sm.cell(row=2, column=13, value=Z).font = Font(name=FONT, color="0000FF")
    srng, lrng = f"key!$B$2:$B${last}", f"Review!${acol}$2:${acol}${last}"
    for i, s in enumerate(STRATA + [CENSUS], start=2):
        put(sm, i, 1, s)
        put(sm, i, 2, f"=COUNTIFS({srng},A{i})")
        put(sm, i, 3, f'=COUNTIFS({srng},A{i},{lrng},"<>")')
        for j, lab in enumerate(LABELS, start=4):
            put(sm, i, j, f'=COUNTIFS({srng},A{i},{lrng},"{lab}")')
        n = f"(C{i}-G{i})"
        put(sm, i, 8, f'=IFERROR(D{i}/{n},"")', numfmt="0%")
        put(sm, i, 9, f'=IFERROR((D{i}+E{i})/{n},"")', numfmt="0%")
        core = f"(H{i}+$M$2^2/(2*{n}))/(1+$M$2^2/{n})"
        half = f"$M$2*SQRT(H{i}*(1-H{i})/{n}+$M$2^2/(4*{n}^2))/(1+$M$2^2/{n})"
        put(sm, i, 10, f'=IFERROR(MAX(0,{core}-{half}),"")', numfmt="0%")
        put(sm, i, 11, f'=IFERROR(MIN(1,{core}+{half}),"")', numfmt="0%")
    nr = len(STRATA) + 4
    for k, t in enumerate([
            "Per stratum, unweighted; 'unclear' is excluded. Census rows count once each here; "
            "the review step weights them by their occurrences.",
            "Weighted figures per source and level, the EU-Dutch difference, lost actors, error "
            "kinds and inventory additions: python actor_map_review.py review --xlsx <this file>"]):
        put(sm, nr + k, 1, t, italic=True)
    for col, w in zip("ABCDEFGHIJK", [22, 7, 8, 8, 8, 8, 8, 8, 8, 13, 13]):
        sm.column_dimensions[col].width = w

    wb.active = wb.sheetnames.index("Review")
    wb.save(path)


# ---------------------------------------------------------------------------
# Review
# ---------------------------------------------------------------------------

def review(args):
    wb = load_workbook(args.xlsx, data_only=True)
    rv = wb["Review"]
    names = [c.value for c in rv[1]]
    ia, ip, ic = (names.index(h) for h in ("Assessment", "What is wrong?", "Correct actor"))
    lab = {}
    for row in rv.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            lab[int(row[0])] = (str(row[ia]).strip().lower() if row[ia] else None,
                                str(row[ip]).strip() if row[ip] else None,
                                str(row[ic]).strip() if row[ic] else "")
    kh = [c.value for c in wb["key"][1]]
    key = {int(r[0]): dict(zip(kh, r)) for r in wb["key"].iter_rows(min_row=2, values_only=True)
           if r[0] is not None}
    N = {r[0]: int(r[3]) for r in wb["weights"].iter_rows(min_row=2, values_only=True) if r[0]}

    judged = {i: v for i, v in lab.items() if v[0] in LABELS}
    bad_labels = sorted(i for i, v in lab.items() if v[0] and v[0] not in LABELS)
    print(f"Rows {len(lab)}   judged {len(judged)}   unclear "
          f"{sum(v[0] == 'unclear' for v in judged.values())}   "
          f"blank {sum(1 for v in lab.values() if not v[0])}")
    if bad_labels:
        print(f"Ignored rows with an unknown assessment: {bad_labels}")
    open_census = [i for i, k in key.items() if k["kind"] == "census"
                   and (lab.get(i, (None,))[0] not in LABELS)]
    if open_census:
        print(f"Note: {len(open_census)} census rows are not judged; their occurrences are "
              "left out of the weighted figures.")
    print()
    if not judged:
        print("Nothing to review yet - fill in 'Assessment' first.")
        return

    rows = [(key[i], v) for i, v in judged.items() if v[0] != "unclear"]
    strict = lambda v: v[0] == "correct"
    lenient = lambda v: v[0] in ("correct", "partial")

    def est(srcs, levels, metric):
        parts = []
        for s in srcs:
            for l in levels:
                name = f"{s} | {l}"
                xs = [v for k, v in rows if k["stratum"] == name]
                parts.append((N.get(name, 0), sum(metric(v) for v in xs), len(xs)))
        n_ex = k_ex = items = 0
        for k, v in rows:
            if k["kind"] == "census" and k["level"] in levels:
                w = sum(int(k[f"occ_{s}"] or 0) for s in srcs)
                if w:
                    n_ex += w
                    k_ex += w * metric(v)
                    items += 1
        return estimate(parts, (n_ex, k_ex, items))

    rep = {"rows": len(lab), "judged": len(judged)}

    # --- per stratum
    print("Per stratum, unweighted (95% Wilson):")
    print(f"  {'stratum':22} {'N_h':>8} {'n':>3}  {'strict':24} {'lenient':24}")
    rep["strata"] = {}
    for s in STRATA:
        xs = [v for k, v in rows if k["stratum"] == s]
        if not xs:
            print(f"  {s:22} {N.get(s, 0):>8,} {0:>3}  (no labels)")
            continue
        n = len(xs)
        ks, kl = sum(map(strict, xs)), sum(map(lenient, xs))
        (l1, h1), (l2, h2) = wilson(ks / n, n), wilson(kl / n, n)
        print(f"  {s:22} {N[s]:>8,} {n:>3}  {ks/n:4.0%} [{l1:3.0%}-{h1:4.0%}]          "
              f"{kl/n:4.0%} [{l2:3.0%}-{h2:4.0%}]")
        rep["strata"][s] = {"N": N[s], "n": n, "strict": ks / n, "lenient": kl / n}
    for s in SOURCES:
        r_ = est([s], LEVELS, strict)
        cx = [(k, v) for k, v in rows if k["kind"] == "census" and int(k[f"occ_{s}"] or 0)]
        if cx:
            occ = sum(int(k[f"occ_{s}"]) for k, _ in cx)
            good = sum(int(k[f"occ_{s}"]) for k, v in cx if strict(v))
            okl = sum(int(k[f"occ_{s}"]) for k, v in cx if lenient(v))
            print(f"  {'census, ' + s:22} {occ:>8,} {len(cx):>3}  {good/occ:4.0%} (exact, by "
                  f"occurrence)  {okl/occ:4.0%}")
            rep["strata"][f"census | {s}"] = {"occurrences": occ, "groups": len(cx),
                                             "strict": good / occ, "lenient": okl / occ}

    # --- weighted per source and claim
    print("\nWeighted by occurrences (strata by N_h, census groups by their occurrences):")
    print(f"  {'':34} {'strict':44} lenient")
    rep["weighted"] = {}
    for claim, levels in [("all results", LEVELS)] + list(CLAIM_LEVELS.items()):
        for srcs, tag in ((["EU"], "EU"), (["NL"], "NL"), (SOURCES, "both")):
            a, b = est(srcs, levels, strict), est(srcs, levels, lenient)
            print(f"  {claim + ', ' + tag:34} {fmt(a):44} {fmt(b)}")
            rep["weighted"][f"{claim} | {tag}"] = {"strict": a, "lenient": b}
    print("\n  by match level, both sources:")
    for l in LEVELS:
        a, b = est(SOURCES, [l], strict), est(SOURCES, [l], lenient)
        print(f"  {l:34} {fmt(a):44} {fmt(b)}")
        rep["weighted"][f"level {l}"] = {"strict": a, "lenient": b}

    # --- EU - Dutch
    print("\nEU - Dutch difference (Newcombe), all results:")
    rep["difference"] = {}
    for mname, metric in (("strict", strict), ("lenient", lenient)):
        a, b = est(["EU"], LEVELS, metric), est(["NL"], LEVELS, metric)
        if a and b:
            d = newcombe(a, b)
            print(f"  {mname:9} raw           {d['d']:+6.1%} [{d['lo']:+6.1%} - {d['hi']:+6.1%}]")
            rep["difference"][f"{mname} raw"] = d
        adj = mix_adjusted(est, metric)
        if adj:
            d = newcombe(*adj)
            print(f"  {mname:9} mix-adjusted  {d['d']:+6.1%} [{d['lo']:+6.1%} - {d['hi']:+6.1%}]")
            rep["difference"][f"{mname} mix-adjusted"] = d
    print("  A gap that remains after the mix adjustment points to the mapper; one that "
          "disappears comes from Dutch units falling into harder match levels.")

    # --- lost actors and missed entries
    print("\nActors lost as 'not an actor' (labelled wrong or partial):")
    rep["lost"] = {}
    for srcs, tag in ((["EU"], "EU"), (["NL"], "NL"), (SOURCES, "both")):
        r_ = est(srcs, ["not_actor"], lambda v: not strict(v))
        if r_:
            print(f"  {tag:5} share {fmt(r_)}   ~{r_['p'] * r_['weight']:,.0f} of "
                  f"{r_['weight']:,.0f} occurrences [{r_['lo'] * r_['weight']:,.0f} - "
                  f"{r_['hi'] * r_['weight']:,.0f}]")
            rep["lost"][tag] = r_
    r_ = est(SOURCES, ["fallback"], lambda v: v[1] == "missed entry")
    if r_:
        print(f"\nFallback units for which an inventory entry fits (missed entry):\n"
              f"  share {fmt(r_)}   ~{r_['p'] * r_['weight']:,.0f} occurrences")
        rep["fallback_missed_entry"] = r_

    # --- error kinds, weighted
    kinds = Counter()
    for k, v in rows:
        if strict(v):
            continue
        if k["kind"] == "census":
            w = int(k["occ_EU"] or 0) + int(k["occ_NL"] or 0)
        else:
            n = sum(1 for kk, _ in rows if kk["stratum"] == k["stratum"])
            w = N.get(k["stratum"], 0) / n if n else 0
        kinds[(CLAIM[k["level"]], v[1] or "unspecified")] += w
    if kinds:
        print("\nError kinds (weighted estimate of occurrences; partial and wrong):")
        for (claim, kind), w in kinds.most_common():
            print(f"  ~{w:>8,.0f}  {claim:13} {kind}")
        rep["error_kinds"] = {f"{c} | {k}": w for (c, k), w in kinds.items()}

    # --- census groups not correct
    cbad = sorted(((k, v) for k, v in rows if k["kind"] == "census" and not strict(v)),
                  key=lambda x: -(int(x[0]["group_occ"] or 0)))
    if cbad:
        print("\nCensus groups not correct (by occurrences):")
        for k, v in cbad:
            print(f"  {int(k['group_occ']):>5}  {k['unit']!s:40} -> {k['result']!s:36} "
                  f"{v[0]:8} {v[1] or ''}{(' -> ' + v[2]) if v[2] else ''}")
        rep["census_not_correct"] = [{"unit": k["unit"], "result": k["result"],
                                      "occurrences": k["group_occ"], "label": v[0],
                                      "problem": v[1], "correct": v[2]} for k, v in cbad]

    # --- inventory additions and missed entries named by the labeller
    adds = defaultdict(dict)   # (problem, actor named) -> {(unit, result): occurrences}
    for k, v in rows:
        if v[1] in ("actor, not in inventory", "missed entry"):
            adds[(v[1], v[2] or k["unit"])][(k["unit"], k["result"])] = int(k["group_occ"] or 0)
    if adds:
        print("\nNamed by the labeller (occurrences of the units concerned, each counted once; "
              "candidates for actor_inventory.py):")
        ranked = sorted(adds.items(), key=lambda x: -sum(x[1].values()))
        for (kind, name), grp in ranked:
            units = sorted({u for u, _ in grp})
            print(f"  {sum(grp.values()):>5}  {kind:24} {name}   (units: {'; '.join(units)[:80]})")
        rep["inventory_candidates"] = [{"kind": k, "actor": n, "occurrences": sum(g.values()),
                                        "units": sorted({u for u, _ in g})} for (k, n), g in ranked]

    # --- verdict
    overall = est(SOURCES, LEVELS, strict)
    if overall:
        lo = overall["lo"]
        verdict = "good" if lo >= 0.75 else "fair" if lo >= 0.60 else "weak"
        print(f"\nVerdict (lower bound of strict precision over all results, both sources): "
              f"{lo:.1%} -> {verdict}  (good >= 75%, fair 60-75%, weak < 60%)")
        rep["verdict"] = {"lower_bound": lo, "verdict": verdict}

    out = Path(str(args.xlsx) + ".report.json")
    out.write_text(json.dumps(rep, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"\n-> {out}")


def mix_adjusted(est, metric):
    """Both sources weighted with the same pooled shares of the match levels."""
    per = {s: {l: est([s], [l], metric) for l in LEVELS} for s in SOURCES}
    levels = [l for l in LEVELS if all(per[s][l] for s in SOURCES)]
    if not levels:
        return None
    w = {l: sum(per[s][l]["weight"] for s in SOURCES) for l in levels}
    tot = sum(w.values())
    out = []
    for s in SOURCES:
        p = sum(w[l] / tot * per[s][l]["p"] for l in levels)
        var = sum((w[l] / tot) ** 2 * per[s][l]["var"] for l in levels)
        n_all = sum(per[s][l]["n"] for l in levels)
        n_eff = min(n_all, p * (1 - p) / var) if var > 0 else n_all
        lo, hi = wilson(p, n_eff)
        out.append({"p": p, "lo": lo, "hi": hi})
    return out


# ---------------------------------------------------------------------------

def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample", help="Write the labelling workbook.")
    s.add_argument("--data", default="data/conflict_all/actor", help="folder of the actor step's output")
    s.add_argument("--norms", default="norm_actor_map.jsonl", help="path, or file name in --data")
    s.add_argument("--units", default="unit_actor_map.jsonl", help="path, or file name in --data")
    s.add_argument("--inventory", default="actor_inventory.jsonl", help="path, or file name in --data")
    s.add_argument("--mapper", default=None, help="path to actor_map.py (default: next to this script)")
    s.add_argument("--max-mismatch", type=float, default=0.01,
                   help="stop if the recheck differs for more than this share of distinct units")
    s.add_argument("--norm-text", default=None, help="norm records with the text (JSON array or JSONL)")
    s.add_argument("--id-field", default="norm_id")
    s.add_argument("--text-field", default="text.paragraph")
    s.add_argument("--text-chars", type=int, default=1500)
    s.add_argument("--census", type=int, default=20, help="most frequent groups judged in full")
    s.add_argument("--contexts", type=int, default=3, help="contexts per census group")
    s.add_argument("--per-source", type=int, default=50, help="sampled units per source")
    s.add_argument("--floor", type=int, default=8, help="minimum per stratum")
    s.add_argument("--alloc", action="append", metavar="STRATUM=N",
                   help='override a stratum size, e.g. --alloc "NL | fallback=12"')
    s.add_argument("--max-per-norm", type=int, default=1)
    s.add_argument("--seed", type=int, default=42)
    s.add_argument("--out", default=None)
    r = sub.add_parser("review", help="Report estimates from a labelled workbook.")
    r.add_argument("--xlsx", default="data/conflict_all/actor/actor_map_review.xlsx")
    args = p.parse_args()
    sample(args) if args.cmd == "sample" else review(args)


if __name__ == "__main__":
    main()