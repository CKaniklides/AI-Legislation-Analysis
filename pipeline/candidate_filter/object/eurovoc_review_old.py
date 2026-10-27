"""
Build an Excel workbook to validate the Step 7 object -> EuroVoc mapping.

One row per unique mapped text (the mapping is computed per text, so judging
per norm would repeat work). With Step 6 these texts are the cleaned object
units; raw_object shows an original object the unit came from. Each row shows the cascade level, up to three
candidate concepts with confidence, how many norms share the object, an
example norm, and a text excerpt around the object for context.

A stratified sample is marked (in_sample = yes) so you can judge a manageable
subset first: --per-stratum objects from each of
  phrase, head, compound, embedding (three percentile bands), unmapped.
Within each stratum the picks are spread over the Step 6 norm types (the
--type-field of the norms behind a text, e.g. unchanged / cleaned / split),
taking turns so that rare types are not missed. Norms that kept no unit at all
(e.g. no_object, placeholder, crossref_only) never reach the mapping; they get
their own strata 'no unit: <type>' with --per-empty-type norms each, to check
that the object was rightly dropped.
The Summary sheet computes precision per stratum with formulas, updating as
you fill in the 'judgement' column.

Sheets
  README    what to fill in, judgement codes, an example row
  Validate  one row per unique object; yellow columns are for you
  Summary   counts and precision per stratum (formulas)
  Norms     every norm with its object and top candidates (reference)

Usage
  python eurovoc_review.py --records <norm records>
  python eurovoc_review.py --records <norm records> --per-stratum 15 --seed 1
  python eurovoc_review.py --records <norm records> --per-stratum 15 --per-empty-type 5
  python eurovoc_review.py --records <norm records> --all-rows   (all texts, sample marked)
Output: <dir>/mapping_review.xlsx
"""

from __future__ import annotations

import argparse
import json
import random
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
JUDGEMENTS = ["correct", "partial", "wrong", "unclear"]
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
HEADER_FILL = PatternFill("solid", start_color="1F3864")
SAMPLE_FILL = PatternFill("solid", start_color="EAF1DD")
THIN = Side(style="thin", color="BFBFBF")


# --------------------------------------------------------------------------
# Inputs
# --------------------------------------------------------------------------

def _get(obj, dotted: str):
    for key in dotted.split("."):
        if not isinstance(obj, dict):
            return None
        obj = obj.get(key)
    return obj


def load_jsonl(path: Path) -> list[dict]:
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def load_records(path: Path) -> list[dict]:
    raw = path.read_text(encoding="utf-8").strip()
    return json.loads(raw) if raw.startswith("[") else \
        [json.loads(l) for l in raw.splitlines() if l.strip()]


def excerpt(paragraph: str | None, obj: str, width: int) -> str:
    """Text around the object in its paragraph; the start of the paragraph if
    the object is not found literally."""
    if not paragraph:
        return ""
    text = " ".join(paragraph.split())
    i = text.lower().find(obj.lower()[:60]) if obj else -1
    if i < 0:
        return text[:2 * width] + ("…" if len(text) > 2 * width else "")
    start, end = max(0, i - width), min(len(text), i + len(obj) + width)
    return ("…" if start else "") + text[start:end] + ("…" if end < len(text) else "")


# --------------------------------------------------------------------------
# Strata and sample
# --------------------------------------------------------------------------

def embedding_bands(threshold: float) -> list[tuple[float, float]]:
    step = (100.0 - threshold) / 3
    return [(threshold + i * step, threshold + (i + 1) * step) for i in range(3)]


def stratum_of(r: dict, bands) -> str:
    level = r["level"]
    if level is None:
        return "unmapped"
    if level != "embedding":
        return level
    pct = (r["candidates"] or [{}])[0].get("percentile") or 0.0
    for lo, hi in bands:
        if pct < hi or hi >= 100:
            return f"embedding p{lo:.0f}-{hi:.0f}"
    return f"embedding p{bands[-1][0]:.0f}-{bands[-1][1]:.0f}"


# --------------------------------------------------------------------------
# Workbook helpers
# --------------------------------------------------------------------------

def style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# --------------------------------------------------------------------------

def main() -> None:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dir", type=Path, default=Path("data/conflict/object/eurovoc"))
    p.add_argument("--records", type=Path, required=True)
    p.add_argument("--object-field", default="norm.object")
    p.add_argument("--id-field", default="norm_id")
    p.add_argument("--text-field", default="text.paragraph")
    p.add_argument("--per-stratum", type=int, default=10)
    p.add_argument("--seed", type=int, default=0)
    p.add_argument("--excerpt-chars", type=int, default=150)
    p.add_argument("--out", type=Path, default=None)
    p.add_argument("--type-field", default="level",
                   help="field with the Step 6 norm type in norm_object_mapping.jsonl or --units")
    p.add_argument("--units", type=Path, default=None,
                   help="Step 6 norm_object_units.jsonl, read for the norm type if the "
                        "mapping file does not have it")
    p.add_argument("--per-empty-type", type=int, default=5,
                   help="sampled norms per type for norms without any unit (0 = skip)")
    p.add_argument("--all-rows", action="store_true",
                   help="write every mapped text to Validate, not only the sample")
    args = p.parse_args()
    out = args.out or args.dir / "mapping_review.xlsx"

    objects = load_jsonl(args.dir / "object_mapping.jsonl")
    norms = load_jsonl(args.dir / "norm_object_mapping.jsonl")
    with open(args.dir / "mapping_config.json", encoding="utf-8") as f:
        threshold = float(json.load(f)["fallback_percentile"])
    bands = embedding_bands(threshold)
    strata_order = ["phrase", "head", "compound"] + \
        [f"embedding p{lo:.0f}-{hi:.0f}" for lo, hi in bands] + ["unmapped"]

    # Context per object: how many norms, one example norm id and paragraph
    paragraph_by_id = {}
    for rec in load_records(args.records):
        paragraph_by_id[str(_get(rec, args.id_field))] = _get(rec, args.text_field)
    # Mapped text -> (norm_id, raw object) of every norm using it. With Step 6
    # output a norm lists its units; without it the raw object is the text.
    norms_per_text = defaultdict(list)
    for n in norms:
        texts = [u["text"] for u in n.get("units") or []] or \
            ([n["object"]] if n["object"] else [])
        for t in dict.fromkeys(texts):
            norms_per_text[t].append((str(n["norm_id"]), n["object"]))

    # Step 6 norm type per norm
    type_by_norm = {str(n["norm_id"]): n.get(args.type_field) for n in norms}
    if not any(type_by_norm.values()) and args.units:
        for u in load_jsonl(args.units):
            type_by_norm[str(u["norm_id"])] = u.get(args.type_field)
    if not any(type_by_norm.values()):
        raise SystemExit(
            f"No '{args.type_field}' field found in norm_object_mapping.jsonl"
            f"{' or ' + str(args.units) if args.units else ''}. Pass --units <Step 6 "
            f"norm_object_units.jsonl> or --type-field <name>. "
            f"Fields in the mapping file: {sorted(norms[0]) if norms else []}")
    for n in norms:
        k = str(n["norm_id"])
        type_by_norm[k] = type_by_norm[k] or ("no_object" if not n["object"] else "unknown")
    type_freq = Counter(type_by_norm.values())

    rows = []
    for r in objects:
        uses = norms_per_text.get(r["object"], [])
        ex_id, raw = uses[0] if uses else ("", "")
        types = sorted({type_by_norm[nid] for nid, _ in uses}) or ["unknown"]
        rows.append({**r, "stratum": stratum_of(r, bands), "n_norms": len(uses),
                     "types": types,
                     # the rarest type decides how the text is spread in the sample
                     "ptype": min(types, key=lambda t: (type_freq[t], t)),
                     "raw_object": raw if raw != r["object"] else "",
                     "example_norm": ex_id,
                     "context": excerpt(paragraph_by_id.get(ex_id), raw or r["object"],
                                        args.excerpt_chars)})
    # Norms that kept no unit: one row per norm, raw object shown as object
    empty_types = []
    if args.per_empty_type > 0:
        for n in norms:
            has_unit = bool(n["units"]) if "units" in n else bool(n["object"])
            if has_unit:
                continue
            t = type_by_norm[str(n["norm_id"])]
            nid = str(n["norm_id"])
            rows.append({"object": n["object"] or "", "raw_object": "", "head": None,
                         "level": None, "matched_label": None, "candidates": [],
                         "stratum": f"no unit: {t}", "n_norms": 1, "types": [t], "ptype": t,
                         "example_norm": nid,
                         "context": excerpt(paragraph_by_id.get(nid), n["object"] or "",
                                            args.excerpt_chars)})
        empty_types = sorted({r["types"][0] for r in rows if r["stratum"].startswith("no unit")})
    strata_order += [f"no unit: {t}" for t in empty_types]

    # Stratified sample
    rng = random.Random(args.seed)
    by_stratum = defaultdict(list)
    for i, r in enumerate(rows):
        by_stratum[r["stratum"]].append(i)
    sample = set()
    for s, idx in by_stratum.items():
        want = args.per_empty_type if s.startswith("no unit") else args.per_stratum
        groups = defaultdict(list)
        for i in idx:
            groups[rows[i]["ptype"]].append(i)
        queues = []
        for t in sorted(groups):
            g = groups[t]
            rng.shuffle(g)
            queues.append(g)
        picked = []
        while len(picked) < want and any(queues):   # take turns over the types
            for q in queues:
                if q and len(picked) < want:
                    picked.append(q.pop())
        sample.update(picked)
    order = {s: i for i, s in enumerate(strata_order)}
    rows_sorted = sorted(
        enumerate(rows),
        key=lambda x: (x[0] not in sample, order.get(x[1]["stratum"], 99),
                       -((x[1]["candidates"] or [{}])[0].get("percentile") or 0),
                       x[1]["object"].lower()))
    counts = Counter(r["stratum"] for r in rows)
    if not args.all_rows:
        rows_sorted = [x for x in rows_sorted if x[0] in sample]

    wb = Workbook()

    # ---------------- README ----------------
    ws = wb.active
    ws.title = "README"
    readme = [
        ("Validation of the object → EuroVoc mapping (Step 7)", True),
        ("", False),
        ("What to do", True),
        ("1. Go to the Validate sheet. Rows with in_sample = yes (green) come first: judge those.", False),
        ("2. Fill in the yellow columns only: judgement (drop-down), better_concept (optional), notes.", False),
        ("3. Judge candidate 1 mainly; use candidates 2-3 to decide 'partial'.", False),
        ("4. The Summary sheet updates automatically as you fill in judgements.", False),
        ("", False),
        ("Judgement codes", True),
        ("correct  – candidate 1 is the right concept, or an acceptable close one.", False),
        ("partial  – candidate 1 is in the right area but too broad/narrow, or a lower candidate is right.", False),
        ("wrong    – no candidate fits the object.", False),
        ("unclear  – the object itself cannot be judged (placeholder, fragment, unresolved reference).", False),
        ("For unmapped objects: 'correct' = rightly left unmapped; 'wrong' = a EuroVoc concept clearly exists.", False),
        ("", False),
        ("Columns", True),
        ("stratum     – cascade level; embedding matches are split into three percentile bands.", False),
        ("object      – the text that was mapped (a cleaned unit from Step 6).", False),
        ("raw_object  – the original object it came from, if different (judge the unit in that light).", False),
        ("norm_type   – Step 6 type of the norms that contain this text (unchanged, cleaned, split, …).", False),
        ("no unit: …  – strata of norms that kept no unit; object shows the raw object (empty for no_object).", False),
        ("              Judge: correct = rightly no unit; wrong = a real object was lost.", False),
        ("level       – phrase / head / compound / embedding, or empty if unmapped.", False),
        ("cand_n, conf_n – candidate concept label and its confidence.", False),
        ("pct_1       – percentile of candidate 1 in the Step 4 top-1 distribution (embedding only).", False),
        ("n_norms     – how many norms share this object; context shows text around it in an example norm.", False),
        ("", False),
        ("Example of a filled-in row (not part of the data)", True),
    ]
    for i, (text, bold) in enumerate(readme, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=13 if i == 1 else 11)
    ex_row = len(readme) + 1
    ex_headers = ["object", "level", "cand_1", "conf_1", "judgement", "better_concept", "notes"]
    ex_values = ["ten minste de in lid 6 bepaalde informatie", "head", "informatie", 0.8,
                 "partial", "melding", "concept right but generic; the object refers to lid 6"]
    for j, (h, v) in enumerate(zip(ex_headers, ex_values), 1):
        ws.cell(row=ex_row, column=j, value=h)
        c = ws.cell(row=ex_row + 1, column=j, value=v)
        c.font = Font(name=FONT)
        if h in ("judgement", "better_concept", "notes"):
            c.fill = INPUT_FILL
    style_header(ws, ex_row, len(ex_headers))
    ws.column_dimensions["A"].width = 45
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 18

    # ---------------- Validate ----------------
    ws = wb.create_sheet("Validate")
    headers = ["in_sample", "stratum", "object", "raw_object", "norm_type", "head", "level", "matched_label",
               "cand_1", "conf_1", "pct_1", "cand_2", "conf_2", "cand_3", "conf_3",
               "n_norms", "example_norm", "context",
               "judgement", "better_concept", "notes"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for _, r in rows_sorted:
        cands = (r["candidates"] or []) + [{}] * 3
        ws.append([
            None,  # in_sample, filled in below
            r["stratum"], r["object"], r["raw_object"], ", ".join(r["types"]),
            r["head"], r["level"], r["matched_label"],
            cands[0].get("label"), cands[0].get("confidence"), cands[0].get("percentile"),
            cands[1].get("label"), cands[1].get("confidence"),
            cands[2].get("label"), cands[2].get("confidence"),
            r["n_norms"], r["example_norm"], r["context"], None, None, None])
    n_rows = len(rows_sorted)
    for k, (i, _) in enumerate(rows_sorted, start=2):
        in_s = i in sample
        ws.cell(row=k, column=1, value="yes" if in_s else "no")
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=k, column=c)
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=headers[c - 1] in ("object", "raw_object", "context", "notes"))
            cell.border = Border(bottom=THIN)
            if headers[c - 1] in ("judgement", "better_concept", "notes"):
                cell.fill = INPUT_FILL
            elif in_s:
                cell.fill = SAMPLE_FILL
        for col in ("conf_1", "conf_2", "conf_3"):
            ws.cell(row=k, column=headers.index(col) + 1).number_format = "0.00"
    last = n_rows + 1
    jcol = get_column_letter(headers.index("judgement") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(JUDGEMENTS) + '"',
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="Judgement", error="Choose: " + ", ".join(JUDGEMENTS))
    ws.add_data_validation(dv)
    dv.add(f"{jcol}2:{jcol}{last}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    widths = {"in_sample": 11, "stratum": 18, "object": 40, "raw_object": 40, "norm_type": 16, "head": 16,
              "level": 11, "matched_label": 16, "cand_1": 26, "conf_1": 8, "pct_1": 7,
              "cand_2": 22, "conf_2": 8, "cand_3": 22, "conf_3": 8, "n_norms": 9,
              "example_norm": 18, "context": 60, "judgement": 12, "better_concept": 22,
              "notes": 30}
    set_widths(ws, {get_column_letter(i): widths[h] for i, h in enumerate(headers, 1)})

    # ---------------- Summary ----------------
    ws = wb.create_sheet("Summary")
    scol = f"Validate!$B$2:$B${last}"
    jrng = f"Validate!${jcol}$2:${jcol}${last}"
    srng = f"Validate!$A$2:$A${last}"
    sum_headers = ["stratum", "objects", "in sample", "judged", "correct", "partial",
                   "wrong", "unclear", "precision (strict)", "precision (lenient)"]
    ws.append(sum_headers)
    style_header(ws, 1, len(sum_headers))
    for i, s in enumerate(strata_order, start=2):
        ws.cell(row=i, column=1, value=s)
        ws.cell(row=i, column=2, value=counts.get(s, 0))  # population, not rows in Validate
        ws.cell(row=i, column=3, value=f'=COUNTIFS({scol},A{i},{srng},"yes")')
        ws.cell(row=i, column=4, value=f'=COUNTIFS({scol},A{i},{jrng},"<>")')
        for j, code in enumerate(JUDGEMENTS, start=5):
            ws.cell(row=i, column=j, value=f'=COUNTIFS({scol},A{i},{jrng},"{code}")')
        ws.cell(row=i, column=9, value=f"=IFERROR(E{i}/(D{i}-H{i}),\"\")")
        ws.cell(row=i, column=10, value=f"=IFERROR((E{i}+F{i})/(D{i}-H{i}),\"\")")
    tot = len(strata_order) + 2
    ws.cell(row=tot, column=1, value="total")
    for c in range(2, 9):
        L = get_column_letter(c)
        ws.cell(row=tot, column=c, value=f"=SUM({L}2:{L}{tot - 1})")
    ws.cell(row=tot, column=9, value=f"=IFERROR(E{tot}/(D{tot}-H{tot}),\"\")")
    ws.cell(row=tot, column=10, value=f"=IFERROR((E{tot}+F{tot})/(D{tot}-H{tot}),\"\")")
    for r in range(2, tot + 1):
        for c in range(1, len(sum_headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, bold=(r == tot))
        ws.cell(row=r, column=9).number_format = "0%"
        ws.cell(row=r, column=10).number_format = "0%"
    note_row = tot + 2
    notes = [
        "Precision excludes 'unclear'. Strict = correct only; lenient = correct + partial.",
        "For 'unmapped', precision means: share of objects rightly left unmapped. "
        "For 'no unit' strata: share of norms that rightly have no unit.",
        f"Embedding bands split the range above the fallback percentile ({threshold:g}) into three equal parts.",
        f"'objects' is the full count per stratum; Validate holds {'all of them' if args.all_rows else 'only the sample'}.",
        f"With {args.per_stratum} judged per stratum, precision per stratum is a rough indication "
        f"(about ±{round(98 / args.per_stratum ** 0.5)} points at worst, 95%); the total is much tighter.",
    ]
    for k, t in enumerate(notes):
        ws.cell(row=note_row + k, column=1, value=t).font = Font(name=FONT, italic=True)
    cov_row = note_row + len(notes) + 1
    ws.cell(row=cov_row, column=1, value="norm_type")
    ws.cell(row=cov_row, column=2, value="norms")
    ws.cell(row=cov_row, column=3, value="rows in sample")
    style_header(ws, cov_row, 3)
    for k, t in enumerate(sorted(type_freq, key=lambda t: -type_freq[t]), start=1):
        ws.cell(row=cov_row + k, column=1, value=t).font = Font(name=FONT)
        ws.cell(row=cov_row + k, column=2, value=type_freq[t]).font = Font(name=FONT)
        ws.cell(row=cov_row + k, column=3,
                value=sum(1 for i in sample if t in rows[i]["types"])).font = Font(name=FONT)
    set_widths(ws, {"A": 22, "B": 9, "C": 10, "D": 8, "E": 8, "F": 8, "G": 8, "H": 8,
                    "I": 17, "J": 18})

    # ---------------- Norms ----------------
    ws = wb.create_sheet("Norms")
    nh = ["norm_id", "object", "units", "level", "cand_1", "conf_1", "cand_2", "cand_3", "norm_type"]
    ws.append(nh)
    style_header(ws, 1, len(nh))
    for n in norms:
        c = (n["candidates"] or []) + [{}] * 3
        units = " | ".join(u["text"] for u in n.get("units") or [])
        ws.append([n["norm_id"], n["object"], units, n["level"], c[0].get("label"),
                   c[0].get("confidence"), c[1].get("label"), c[2].get("label"),
                   type_by_norm[str(n["norm_id"])]])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical="top", wrap_text=(cell.column in (2, 3)))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:I{len(norms) + 1}"
    set_widths(ws, {"A": 22, "B": 50, "C": 35, "D": 11, "E": 26, "F": 8, "G": 22, "H": 22,
                    "I": 16})

    wb.active = wb.sheetnames.index("Validate")
    wb.save(out)
    print(f"Wrote {out}")
    print(f"Mapped texts: {len(objects)}, norms without a unit: {len(rows) - len(objects)}, in sample: {len(sample)}, rows in Validate: {len(rows_sorted)}")
    for s in strata_order:
        print(f"  {s:<22} {counts.get(s, 0):>5}")
    print("Norm types (norms / sampled rows containing the type):")
    for t in sorted(type_freq, key=lambda t: -type_freq[t]):
        print(f"  {t:<22} {type_freq[t]:>5} {sum(1 for i in sample if t in rows[i]['types']):>5}")


if __name__ == "__main__":
    main()