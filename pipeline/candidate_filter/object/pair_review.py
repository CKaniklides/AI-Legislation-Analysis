#!/usr/bin/env python3
"""pair_review.py - validation of the object pair scores (object filter).

Two steps:
  1. sample  - draw a stratified sample of scored norm pairs, above and below the
               threshold, and write a workbook for blind labelling.
  2. review  - read the labels back and report candidate precision, the share of
               related pairs at each score level, the related pairs lost below the
               threshold and the direction check, with 95% intervals.

Population
  The Stage A pairs (stageA_pairs.json). A pair's object score is spec x weight of
  the best shared key; pairs that share no key have no score. Pairs without a score
  make no claim and are not sampled: their numbers are exact counts, split into
  pairs where a norm has no object unit and pairs whose norms share no key.

Strata (justified in the workbook's 'Design' sheet)
  Above the threshold (candidates): relation x score level. head and eurovoc are
  split at --upper (default 0.6); noun and tail are too small above the threshold to
  split. Below the threshold: score level only, --lower (default 0.4) to the
  threshold (near misses) and below --lower; --low-split adds a fifth level.

  Default allocation (120 judgements)
    above | head | 0.60+        12     above | eurovoc | 0.60+        10
    above | head | 0.50-0.60    20     above | eurovoc | 0.50-0.60    15
    above | noun | 0.50+         8     above | tail | 0.50+            8
    below | 0.40-0.50           22     below | <0.40                  25
  A norm appears in at most --max-uses sampled pairs (default 2).

Labels: same / related / different / unclear. A candidate is right when judged
same or related (lenient) or same (strict); a pair below the threshold is lost
when judged same or related. For same/related pairs a second question asks which
object is broader, to check the direction flag.

Usage
  python pair_review.py sample --stagea data/stageA_pairs.json --norm-text data/norms_all.json
  python pair_review.py sample ... --low-split 0.35 --alloc "below | 0.35-0.40=15"
  python pair_review.py review --xlsx data/conflict/object/pair_review.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import random
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RELATIONS = ["head", "eurovoc", "tail", "noun"]
SPLIT_RELATIONS = {"head", "eurovoc"}          # split at --upper above the threshold
JUDGEMENTS = ["same", "related", "different", "unclear"]
BROADER = ["A broader", "B broader", "equal scope", "can't say"]
# direction flag -> the answer that agrees with it
FLAG_ANSWER = {"same": "equal scope", "a_narrower": "B broader", "b_narrower": "A broader"}
OK = {"strict": {"same"}, "lenient": {"same", "related"}}
Z = 1.96

FONT = "Arial"
INPUT_FILL = PatternFill("solid", start_color="FFFF00")
HEAD_FILL = PatternFill("solid", start_color="1F3864")
THIN = Side(style="thin", color="BFBFBF")


# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------

def _get(obj, dotted):
    for k in dotted.split("."):
        if not isinstance(obj, dict):
            return None
        obj = obj.get(k)
    return obj


def read_jsonl(path):
    with open(path, encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def read_records(path):
    raw = Path(path).read_text(encoding="utf-8").strip()
    return json.loads(raw) if raw.startswith("[") else \
        [json.loads(l) for l in raw.splitlines() if l.strip()]


def unit_text(u):
    if isinstance(u, str):
        return u
    return u.get("unit") or u.get("text") or ""


def source_of(norm_id):
    return "NL" if str(norm_id).upper().startswith("BWBR") else "EU"


def load_stagea(path):
    """(norm_ids, pairs as (i, j)) from stageA_filtering.py output."""
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    if not (isinstance(data, dict) and "norm_ids" in data and "pairs" in data):
        raise SystemExit(f"{path}: expected the Stage A format with 'norm_ids' and 'pairs'.")
    fmt = list(data.get("pair_format") or ["pair_id", "i", "j"])
    ci, cj = fmt.index("i"), fmt.index("j")
    return [str(x) for x in data["norm_ids"]], [(p[ci], p[cj]) for p in data["pairs"]]


# ---------------------------------------------------------------------------
# Strata
# ---------------------------------------------------------------------------

def f2(x):
    return f"{x:.2f}"


class Design:
    """Stratum and score-level definitions from the thresholds."""

    def __init__(self, threshold, upper, lower, low_split):
        if not lower < threshold < upper:
            raise SystemExit("Need --lower < --threshold < --upper.")
        if low_split is not None and not low_split < lower:
            raise SystemExit("--low-split must be below --lower.")
        self.t, self.up, self.lo, self.ls = threshold, upper, lower, low_split
        # score levels, highest first
        self.levels = [f"{f2(upper)}+", f"{f2(threshold)}-{f2(upper)}", f"{f2(lower)}-{f2(threshold)}"]
        self.levels += ([f"{f2(low_split)}-{f2(lower)}", f"<{f2(low_split)}"] if low_split is not None
                        else [f"<{f2(lower)}"])

    def level(self, s):
        if s >= self.up:
            return self.levels[0]
        if s >= self.t:
            return self.levels[1]
        if s >= self.lo:
            return self.levels[2]
        if self.ls is not None and s >= self.ls:
            return self.levels[3]
        return self.levels[-1]

    def stratum(self, p):
        s, r = p["score"], p["relation"]
        if s >= self.t:
            if r in SPLIT_RELATIONS:
                return f"above | {r} | {self.level(s)}"
            return f"above | {r} | {f2(self.t)}+"
        return f"below | {self.level(s)}"

    def default_n(self, name):
        if name.startswith("below"):
            lvl = name.split(" | ")[1]
            if lvl == self.levels[2]:
                return 22
            if self.ls is None:
                return 25
            return 13 if lvl == self.levels[3] else 12
        _, rel, lvl = name.split(" | ")
        if rel in SPLIT_RELATIONS:
            high = lvl == self.levels[0]
            return {"head": (20, 12), "eurovoc": (15, 10)}[rel][high]
        return 8

    def why(self, name):
        if name.startswith("below"):
            lvl = name.split(" | ")[1]
            if lvl == self.levels[2]:
                return ("Near misses just below the threshold: if as many are related as just above "
                        "it, the threshold could come down.")
            if self.ls is not None and lvl == self.levels[3]:
                return "Mostly head and eurovoc pairs sharing generic keys."
            return ("The bulk of scored pairs, mostly generic keys and low-weight noun/tail keys: "
                    "how many related pairs are lost.")
        _, rel, lvl = name.split(" | ")
        if rel in SPLIT_RELATIONS:
            if lvl == self.levels[0]:
                return f"Does a higher score mean higher precision? ({rel}, weight claim)"
            return (f"{'Largest' if rel == 'head' else 'Second-largest'} candidate cell, right at the "
                    "threshold." + (" Also depends on the EuroVoc mapping." if rel == "eurovoc" else ""))
        return f"{rel} keys (lower weight); too few above the threshold to split; floor only."


# ---------------------------------------------------------------------------
# Statistics
# ---------------------------------------------------------------------------

def wilson(p, n, z=Z):
    if n <= 0:
        return 0.0, 1.0
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(max(0.0, p * (1 - p) / n + z * z / (4 * n * n))) / d
    return max(0.0, c - h), min(1.0, c + h)


def combine(parts):
    """parts: (weight N, k, n, population of the stratum). Stratified proportion with a
    Wilson interval on the effective sample size (capped at the labelled count)."""
    parts = [x for x in parts if x[0] > 0 and x[2] > 0]
    if not parts:
        return None
    total = sum(w for w, *_ in parts)
    p = var = 0.0
    n_all = 0
    for w, k, n, pop in parts:
        ph = k / n
        fpc = max(0.0, 1 - n / pop) if pop else 1.0
        p += w / total * ph
        var += (w / total) ** 2 * fpc * ph * (1 - ph) / n
        n_all += n
    n_eff = min(n_all, p * (1 - p) / var) if var > 0 else n_all
    lo, hi = wilson(p, n_eff)
    return p, lo, hi, n_eff


def fmt(res):
    if res is None:
        return "(no labels)"
    p, lo, hi, n_eff = res
    return f"{p:5.1%} [{lo:5.1%} - {hi:5.1%}]   n_eff {n_eff:5.1f}"


# ---------------------------------------------------------------------------
# Sampling
# ---------------------------------------------------------------------------

def sample(args):
    data = Path(args.data)
    out = Path(args.out) if args.out else data / "pair_review.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    design = Design(args.threshold, args.upper, args.lower, args.low_split)

    def resolve(name):
        """A path as given (relative to the working folder) if it exists, else inside --data."""
        p = Path(name)
        return p if p.is_file() else data / p

    norms_path = resolve(args.norms)
    if not norms_path.is_file():
        raise SystemExit(f"--norms: {norms_path} not found. It must be the object step's per-norm "
                         "file (norm_id, object, units), e.g. norm_keys.jsonl in --data.")
    records = read_records(norms_path)
    if records and not any("units" in r for r in records[:200]):
        raise SystemExit(
            f"--norms: {norms_path} has no 'units' field, so it is not the per-norm object file. "
            "Pass the object step's norm_keys.jsonl (norm_id, object, units) as --norms; norm "
            "records with the paragraph text (e.g. norms_all.json) go to --norm-text.")
    norms = {}
    for r in records:
        units = [t for t in (unit_text(u) for u in r.get("units") or []) if t]
        norms[str(r["norm_id"])] = {"object": r.get("object") or "", "units": units}
    pairs_path = resolve(args.pairs)
    if not pairs_path.is_file():
        raise SystemExit(f"--pairs: {pairs_path} not found.")
    pairs = read_jsonl(pairs_path)
    texts = {}
    if args.norm_text:
        for r in read_records(args.norm_text):
            texts[str(_get(r, args.id_field))] = str(_get(r, args.text_field) or "")[:args.text_chars]
    known = []
    if args.known:
        with open(args.known, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                known.append(tuple(sorted((str(r["norm_a"]).strip(), str(r["norm_b"]).strip()))))

    # --- population check and exact counts against Stage A -------------------
    counts = {"scored": len(pairs)}
    if args.stagea:
        norm_ids, spairs = load_stagea(args.stagea)
        pos = {nid: k for k, nid in enumerate(norm_ids)}
        scored_ij = set()
        outside = 0
        for p in pairs:
            ia, ib = pos.get(p["norm_a"]), pos.get(p["norm_b"])
            if ia is None or ib is None:
                outside += 1
                continue
            scored_ij.add((min(ia, ib), max(ia, ib)))
        has_unit = [bool(norms.get(nid, {}).get("units")) for nid in norm_ids]
        in_stagea = empty = 0
        for i, j in spairs:
            if (i, j) in scored_ij:
                in_stagea += 1
            elif not (has_unit[i] and has_unit[j]):
                empty += 1
        outside += len(scored_ij) - in_stagea
        if outside:
            msg = (f"{outside:,} scored pair(s) are not Stage A pairs. The scored file and "
                   f"{args.stagea} do not describe the same population.")
            if not args.allow_outside:
                raise SystemExit(msg + " Fix the inputs, or pass --allow-outside to sample anyway.")
            print("Warning: " + msg)
        counts.update(stagea=len(spairs), empty=empty,
                      no_key=len(spairs) - in_stagea - empty, outside=outside)
    else:
        print("Warning: no --stagea given; the population is not checked and the unscored "
              "pairs are not counted.")

    # --- strata ---------------------------------------------------------------
    pools = defaultdict(list)
    for p in pairs:
        pools[design.stratum(p)].append(p)
    def order(s):
        parts = s.split(" | ")
        rel = parts[1] if parts[0] == "above" else ""
        lvl = parts[-1]
        return (parts[0] == "below", RELATIONS.index(rel) if rel in RELATIONS else len(RELATIONS),
                design.levels.index(lvl) if lvl in design.levels else -1)
    names = sorted(pools, key=order)
    alloc = {s: design.default_n(s) for s in names}
    for item in args.alloc or []:
        name, _, n = item.rpartition("=")
        if name.strip() not in alloc:
            raise SystemExit(f"--alloc: unknown stratum '{name.strip()}'. Strata: {list(alloc)}")
        alloc[name.strip()] = int(n)
    unknown_rel = {p["relation"] for p in pairs} - set(RELATIONS)
    if unknown_rel:
        print(f"Warning: unexpected relation(s) {sorted(unknown_rel)}; given a floor of 8 when above the threshold.")
    N = {s: len(pools[s]) for s in names}
    N_level = {}
    for s in names:
        for lvl, c in Counter(design.level(p["score"]) for p in pools[s]).items():
            N_level[(s, lvl)] = c

    # top keys among candidates (for the key-frequency breakdown)
    cand_keys = Counter(p["key"] for p in pairs if p["score"] >= design.t)
    top_keys = {k for k, _ in cand_keys.most_common(args.top_keys)}

    # --- draw: smallest strata first so the per-norm cap does not starve them ---
    rng = random.Random(args.seed)
    uses = Counter()
    rows = []
    for s in sorted(names, key=lambda s: N[s]):
        pool = sorted(pools[s], key=lambda p: (p["norm_a"], p["norm_b"]))
        rng.shuffle(pool)
        taken = 0
        for p in pool:
            if taken == alloc[s]:
                break
            if uses[p["norm_a"]] >= args.max_uses or uses[p["norm_b"]] >= args.max_uses:
                continue
            uses[p["norm_a"]] += 1
            uses[p["norm_b"]] += 1
            taken += 1
            rows.append(dict(p, stratum=s, level=design.level(p["score"])))
        if taken < alloc[s]:
            print(f"Warning: '{s}' yielded {taken} pairs (target {alloc[s]}).")
    rng.shuffle(rows)   # order must not reveal the score

    info = dict(design=design, names=names, N=N, N_level=N_level, alloc=alloc, counts=counts,
                taken=Counter(r["stratum"] for r in rows), top_keys=top_keys,
                cand_total=sum(N[s] for s in names if s.startswith("above")),
                below_total=sum(N[s] for s in names if s.startswith("below")), args=args)
    write_workbook(out, rows, norms, texts, known, {(p["norm_a"], p["norm_b"]): p for p in pairs},
                   info)

    print(f"Wrote {len(rows)} pairs -> {out}")
    if "stagea" in counts:
        print(f"Stage A pairs {counts['stagea']:,}: scored {counts['scored']:,}, "
              f"a norm without object unit {counts['empty']:,}, units but no shared key {counts['no_key']:,}")
    print(f"{'stratum':32} {'N_h':>8} {'n':>4}")
    for s in names:
        print(f"{s:32} {N[s]:>8,} {info['taken'][s]:>4}")


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

def header(ws, row, ncols, col0=1):
    for c in range(col0, col0 + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def put(ws, row, col, value, wrap=False, bold=False, italic=False, numfmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, bold=bold, italic=italic)
    c.alignment = Alignment(vertical="top", wrap_text=wrap)
    if numfmt:
        c.number_format = numfmt
    return c


def expected_widths(info):
    """Worst-case (p = 0.5) 95% half-widths, in points."""
    N, alloc = info["N"], info["alloc"]

    def hw(prefix):
        names = [s for s in info["names"] if s.startswith(prefix) and alloc[s]]
        tot = sum(N[s] for s in names)
        v = sum((N[s] / tot) ** 2 * max(0.0, 1 - alloc[s] / N[s]) * 0.25 / alloc[s] for s in names)
        return 100 * Z * math.sqrt(v)
    return hw("above"), hw("below")


def design_note(info):
    d, c, a = info["design"], info["counts"], info["args"]
    hw_c, hw_b = expected_widths(info)
    pop = (f"The population is the {c['stagea']:,} Stage A pairs. Of these, {c['scored']:,} share at "
           f"least one key and have an object score; {c['empty']:,} involve a norm without any object "
           f"unit and {c['no_key']:,} involve two norms whose units share no key. "
           if "stagea" in c else
           f"The population is the {c['scored']:,} scored pairs (Stage A was not checked). ")
    return (
        "What is validated. The object score of a pair is the specificity of the best key the two "
        "norms' objects share, times a weight for the kind of match (head 1.0, eurovoc 0.9, tail 0.8, "
        f"noun 0.7). Pairs scoring {d.t:g} or more are passed on as candidates. The question for the "
        "labeller is whether the two objects are the same thing, related, or different.\n\n"
        "Population. " + pop +
        "Pairs without a score make no claim, so they are not sampled; their numbers are exact "
        "counts (table below). Related pairs among them cannot be measured here.\n\n"
        "Why these strata. The filter makes two claims: a pair at or above the threshold is related "
        "(precision), a scored pair below it is not related enough (pairs lost). Both are sampled. "
        "Score levels on each side of the threshold show whether the share of related pairs changes "
        f"at the cut-off: {', '.join(d.levels)}. Two levels per side are the minimum to see a trend; "
        "more would leave too few judgements per level to see one. Above the threshold the levels "
        "are crossed with the relation of the deciding key, because candidates are what the pipeline "
        "uses and each relation's weight is a claim to test; noun and tail are not split, having "
        f"almost no pairs at {d.up:g} or more. Below the threshold the levels alone are enough.\n\n"
        "Why these sizes. Candidates receive most judgements although they are a minority of scored "
        "pairs, because candidate precision is the main figure and the candidate cells are small. "
        "Every cell gets a floor; the rest goes to the largest cells. A norm appears in at most "
        f"{a.max_uses} sampled pairs. At 50% in every cell this gives about +/-{hw_c:.0f} points on "
        f"candidate precision and +/-{hw_b:.0f} points on the share of related pairs below the "
        "threshold. Overall figures are weighted by N_h (the 'review' step), never by row count.\n\n"
        "Caveat. Keys are concentrated: pairs sharing a frequent key tend to be right or wrong "
        "together, so the intervals, which treat pairs as independent, are slightly optimistic."
    )


def write_workbook(path, rows, norms, texts, known, by_pair, info):
    d, N, a = info["design"], info["N"], info["args"]
    wb = Workbook()

    # ---------------- Instructions ----------------
    g = wb.active
    g.title = "Instructions"
    lines = [
        ("How to judge", True),
        ("Fill in the yellow cells of the Review sheet, one row per pair. Judge only the objects.", False),
        ("Same thing?   same: both objects are the same thing.", False),
        ("              related: overlapping, or one is a part, kind or case of the other.", False),
        ("              different: not about the same thing.   unclear: cannot tell (not counted).", False),
        ("Broader?      only when same or related: which object is broader?", False),
        ("              A broader / B broader / equal scope / can't say.", False),
        ("Units show how each object was split; judge the objects as a whole.", False),
        ("Scores are in the hidden 'key' sheet: leave it alone until you are done.", False),
        ("When finished, save and run: python pair_review.py review --xlsx <this file>", False),
        ("", False),
        ("Example (format only; not part of the sample)", True),
        ("Object A: de risicobeoordeling | Object B: een beoordeling van de risico's", False),
        ("Same thing? same | Broader? equal scope | Note: same assessment, other wording", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        put(g, i, 1, text, bold=bold)
    g.column_dimensions["A"].width = 100

    # ---------------- Review (blind) ----------------
    ws = wb.create_sheet("Review")
    cols = [("#", 5), ("Norm A", 22), ("Object A", 40), ("Units A", 30)]
    if texts:
        cols += [("Text A", 55)]
    cols += [("Norm B", 22), ("Object B", 40), ("Units B", 30)]
    if texts:
        cols += [("Text B", 55)]
    cols += [("Same thing?", 13), ("Broader?", 13), ("Note", 30)]
    names = [c for c, _ in cols]
    for j, (name, width) in enumerate(cols, 1):
        ws.cell(row=1, column=j, value=name)
        ws.column_dimensions[get_column_letter(j)].width = width
    header(ws, 1, len(cols))
    for r, p in enumerate(rows, 2):
        na, nb = norms.get(p["norm_a"], {}), norms.get(p["norm_b"], {})
        values = {"#": r - 1, "Norm A": p["norm_a"], "Object A": na.get("object", ""),
                  "Units A": " | ".join(na.get("units", [])), "Text A": texts.get(p["norm_a"], ""),
                  "Norm B": p["norm_b"], "Object B": nb.get("object", ""),
                  "Units B": " | ".join(nb.get("units", [])), "Text B": texts.get(p["norm_b"], "")}
        for j, name in enumerate(names, 1):
            c = put(ws, r, j, values.get(name), wrap=name not in ("#",))
            c.border = Border(bottom=THIN)
            if name in ("Same thing?", "Broader?", "Note"):
                c.fill = INPUT_FILL
    last = len(rows) + 1
    jcol = get_column_letter(names.index("Same thing?") + 1)
    bcol = get_column_letter(names.index("Broader?") + 1)
    for col, options in ((jcol, JUDGEMENTS), (bcol, BROADER)):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
    ws.freeze_panes = "B2"

    # ---------------- key (hidden) ----------------
    ks = wb.create_sheet("key")
    kh = ["#", "stratum", "level", "score", "relation", "key", "spec", "weight", "direction",
          "top_key", "sources", "unit_a", "unit_b"]
    ks.append(kh)
    for i, p in enumerate(rows, 1):
        ks.append([i, p["stratum"], p["level"], p["score"], p["relation"], p["key"], p.get("spec"),
                   p.get("weight"), p.get("direction") or "", "yes" if p["key"] in info["top_keys"] else "no",
                   "-".join(sorted((source_of(p["norm_a"]), source_of(p["norm_b"])))),
                   p.get("unit_a"), p.get("unit_b")])
    for row in ks.iter_rows():
        for c in row:
            c.font = Font(name=FONT)
    ks.sheet_state = "hidden"

    # ---------------- levels (hidden; population per stratum x level) ----------------
    lv = wb.create_sheet("levels")
    lv.append(["stratum", "level", "N"])
    for lvl in d.levels:                      # highest level first; review keeps this order
        for s in info["names"]:
            if info["N_level"].get((s, lvl)):
                lv.append([s, lvl, info["N_level"][(s, lvl)]])
    lv.sheet_state = "hidden"

    # ---------------- Design ----------------
    ds = wb.create_sheet("Design")
    put(ds, 1, 1, "Sample design", bold=True)
    ds.merge_cells("A2:G2")
    put(ds, 2, 1, design_note(info), wrap=True)
    ds.row_dimensions[2].height = 380
    th = ["stratum", "side", "relation", "population N_h", "share of scored pairs", "n sampled",
          "why this stratum and this n"]
    hr = 4
    for j, h in enumerate(th, 1):
        ds.cell(row=hr, column=j, value=h)
    header(ds, hr, len(th))
    first, lastrow = hr + 1, hr + len(info["names"])
    for k, s in enumerate(info["names"]):
        r = first + k
        parts = s.split(" | ")
        put(ds, r, 1, s)
        put(ds, r, 2, parts[0])
        put(ds, r, 3, parts[1] if parts[0] == "above" else "all")
        put(ds, r, 4, N[s], numfmt="#,##0")
        put(ds, r, 5, f"=D{r}/SUM($D${first}:$D${lastrow})", numfmt="0.0%")
        put(ds, r, 6, info["taken"][s])
        put(ds, r, 7, d.why(s), wrap=True)
    tr = lastrow + 1
    put(ds, tr, 1, "total", bold=True)
    put(ds, tr, 4, f"=SUM(D{first}:D{lastrow})", bold=True, numfmt="#,##0")
    put(ds, tr, 6, f"=SUM(F{first}:F{lastrow})", bold=True)

    c = info["counts"]
    r = tr + 2
    put(ds, r, 1, "Exact counts (not sampled)", bold=True)
    ds.cell(row=r + 1, column=1, value="pairs")
    ds.cell(row=r + 1, column=2, value="count")
    ds.cell(row=r + 1, column=3, value="share of Stage A")
    header(ds, r + 1, 3)
    exact = [("Stage A pairs", c.get("stagea")), ("scored (sampled above)", c["scored"]),
             ("a norm without any object unit", c.get("empty")),
             ("both norms have units, no shared key", c.get("no_key"))]
    for k, (label, v) in enumerate(exact):
        rr = r + 2 + k
        put(ds, rr, 1, label)
        put(ds, rr, 2, v if v is not None else "not checked", numfmt="#,##0")
        if c.get("stagea"):
            put(ds, rr, 3, f"=B{rr}/$B${r + 2}", numfmt="0.0%")
    rr = r + 2 + len(exact) + 1
    put(ds, rr, 1, f"Settings: threshold {d.t:g}, upper {d.up:g}, lower {d.lo:g}"
                   f"{', low split ' + format(d.ls, 'g') if d.ls is not None else ''}, "
                   f"max uses per norm {a.max_uses}, top keys {a.top_keys}, seed {a.seed}", italic=True)
    for col, w in zip("ABCDEFG", [34, 10, 16, 14, 12, 10, 70]):
        ds.column_dimensions[col].width = w

    # ---------------- Summary (live, unweighted per stratum) ----------------
    sm = wb.create_sheet("Summary")
    sh = ["stratum", "population N_h", "rows", "judged", "same", "related", "different", "unclear",
          "strict", "lenient", "lenient 95% low", "lenient 95% high"]
    sm.append(sh)
    header(sm, 1, len(sh))
    sm.cell(row=1, column=14, value="z")
    sm.cell(row=2, column=14, value=Z)
    srng, jrng = f"key!$B$2:$B${last}", f"Review!${jcol}$2:${jcol}${last}"
    for i, s in enumerate(info["names"], start=2):
        put(sm, i, 1, s)
        put(sm, i, 2, N[s], numfmt="#,##0")
        put(sm, i, 3, f"=COUNTIFS({srng},A{i})")
        put(sm, i, 4, f'=COUNTIFS({srng},A{i},{jrng},"<>")')
        for j, code in enumerate(JUDGEMENTS, start=5):
            put(sm, i, j, f'=COUNTIFS({srng},A{i},{jrng},"{code}")')
        n = f"(D{i}-H{i})"
        put(sm, i, 9, f'=IFERROR(E{i}/{n},"")', numfmt="0%")
        put(sm, i, 10, f'=IFERROR((E{i}+F{i})/{n},"")', numfmt="0%")
        core = f"(J{i}+$N$2^2/(2*{n}))/(1+$N$2^2/{n})"
        half = f"$N$2*SQRT(J{i}*(1-J{i})/{n}+$N$2^2/(4*{n}^2))/(1+$N$2^2/{n})"
        put(sm, i, 11, f'=IFERROR(MAX(0,{core}-{half}),"")', numfmt="0%")
        put(sm, i, 12, f'=IFERROR(MIN(1,{core}+{half}),"")', numfmt="0%")
    nr = len(info["names"]) + 3
    for k, t in enumerate([
            "Per stratum, unweighted. Above the threshold 'lenient' is precision; below it, the share of "
            "related pairs that were lost.",
            "Weighted figures (candidate precision, score-level curve, pairs lost, direction): "
            "python pair_review.py review --xlsx <this file>"]):
        put(sm, nr + k, 1, t, italic=True)
    for col, w in zip("ABCDEFGHIJKL", [32, 13, 7, 8, 7, 8, 9, 8, 8, 8, 14, 14]):
        sm.column_dimensions[col].width = w

    # ---------------- Known (diagnostic only) ----------------
    if known:
        kn = wb.create_sheet("Known")
        kn.append(["norm_a", "norm_b", "score", "relation", "level", "candidate?"])
        header(kn, 1, 6)
        for a_, b_ in known:
            p = by_pair.get((a_, b_))
            kn.append([a_, b_, p["score"] if p else 0.0, p["relation"] if p else "",
                       d.level(p["score"]) if p else "no score",
                       "yes" if p and p["score"] >= d.t else "no"])
        put(kn, len(known) + 3, 1, "Diagnostic only: known pairs are not a random sample and "
                                  "never enter the estimates.", italic=True)

    wb.active = wb.sheetnames.index("Review")
    wb.save(path)


# ---------------------------------------------------------------------------
# Review
# ---------------------------------------------------------------------------

def review(args):
    wb = load_workbook(args.xlsx, data_only=True)
    rv = wb["Review"]
    names = [c.value for c in rv[1]]
    jc, bc = names.index("Same thing?"), names.index("Broader?")
    lab = {}
    for row in rv.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            lab[int(row[0])] = (str(row[jc]).strip().lower() if row[jc] else None,
                                str(row[bc]).strip() if row[bc] else None)
    kh = [c.value for c in wb["key"][1]]
    key = {int(r[0]): dict(zip(kh, r)) for r in wb["key"].iter_rows(min_row=2, values_only=True)
           if r[0] is not None}
    N = {}
    for row in wb["Design"].iter_rows(values_only=True):
        if row and isinstance(row[0], str) and row[0].startswith(("above |", "below |")) \
                and isinstance(row[3], (int, float)):
            N[row[0]] = int(row[3])
    N_level = {(r[0], r[1]): int(r[2]) for r in wb["levels"].iter_rows(min_row=2, values_only=True)
               if r[0]}
    levels = list(dict.fromkeys(l for (_, l) in N_level))   # sheet order: highest level first

    judged = {i: v for i, v in lab.items() if v[0] in JUDGEMENTS}
    print(f"Rows {len(lab)}   judged {len(judged)}   unclear "
          f"{sum(v[0] == 'unclear' for v in judged.values())}   blank {sum(1 for v in lab.values() if not v[0])}\n")
    if not judged:
        print("Nothing to review yet - fill in 'Same thing?' first.")
        return

    def rows_of(pred):
        return [(key[i], v) for i, v in judged.items() if v[0] != "unclear" and pred(key[i])]

    def part(stratum, mode, pred=lambda k: True, weight=None):
        xs = rows_of(lambda k: k["stratum"] == stratum and pred(k))
        k = sum(v[0] in OK[mode] for _, v in xs)
        return (N[stratum] if weight is None else weight, k, len(xs), N[stratum])

    above = [s for s in N if s.startswith("above")]
    below = [s for s in N if s.startswith("below")]

    print("Per stratum (95% Wilson), share judged same/related (lenient) and same (strict):")
    print(f"  {'stratum':32} {'N_h':>8} {'n':>3}  {'lenient':15} {'strict':15}")
    for s in above + below:
        _, kl, n, _ = part(s, "lenient")
        _, ks_, _, _ = part(s, "strict")
        if n:
            lo, hi = wilson(kl / n, n)
            lo2, hi2 = wilson(ks_ / n, n)
            print(f"  {s:32} {N[s]:>8,} {n:>3}  {kl/n:4.0%} [{lo:3.0%}-{hi:4.0%}]  "
                  f"{ks_/n:4.0%} [{lo2:3.0%}-{hi2:4.0%}]")
        else:
            print(f"  {s:32} {N[s]:>8,} {0:>3}  (no labels)")

    print("\nCandidate precision (score >= threshold), weighted by N_h:")
    for mode in ("lenient", "strict"):
        print(f"  {mode:8} {fmt(combine([part(s, mode) for s in above]))}")

    print("\nShare judged same or related per score level (weighted; the threshold sits between "
          "the second and third level):")
    for lvl in levels:
        parts = [part(s, "lenient", lambda k, l=lvl: k["level"] == l, weight=N_level[(s, lvl)])
                 for s in N if N_level.get((s, lvl))]
        missing = [s for s in N if N_level.get((s, lvl)) and not part(
            s, "lenient", lambda k, l=lvl: k["level"] == l)[2]]
        note = ("   (no labels for " + ", ".join(f"{s} [{N_level[(s, lvl)]:,}]" for s in missing) + ")") \
            if missing else ""
        print(f"  {lvl:12} {fmt(combine(parts))}{note}")

    res = combine([part(s, "lenient") for s in below])
    print("\nRelated pairs lost below the threshold (judged same or related):")
    if res:
        tot = sum(N[s] for s in below)
        print(f"  share    {fmt(res)}")
        print(f"  pairs    ~{res[0] * tot:,.0f} of {tot:,} [{res[1] * tot:,.0f} - {res[2] * tot:,.0f}]")

    # threshold decision over all scored pairs: right when above & related, or below & different
    parts = []
    for s in above + below:
        xs = rows_of(lambda k, s=s: k["stratum"] == s)
        good = sum((v[0] in OK["lenient"]) == s.startswith("above") for _, v in xs)
        parts.append((N[s], good, len(xs), N[s]))
    print(f"\nThreshold decision correct, all scored pairs:  {fmt(combine(parts))}")

    print("\nCandidate precision per relation (lenient, levels pooled by N_h):")
    for rel in RELATIONS:
        ss = [s for s in above if s.split(" | ")[1] == rel]
        if ss:
            print(f"  {rel:8} {fmt(combine([part(s, 'lenient') for s in ss]))}")

    print("\nCandidate precision by group (lenient; descriptive, groups too small to compare):")
    for title, field, values in (("key", "top_key", ("yes", "no")),
                                 ("sources", "sources", ("EU-EU", "EU-NL", "NL-NL"))):
        for val in values:
            xs = rows_of(lambda k, f=field, v=val: k["stratum"] in above and k[f] == v)
            if xs:
                k = sum(v[0] in OK["lenient"] for _, v in xs)
                lo, hi = wilson(k / len(xs), len(xs))
                label = {"yes": "top keys", "no": "other keys"}.get(val, val)
                print(f"  {label:12} {k}/{len(xs)} = {k/len(xs):.0%} [{lo:.0%} - {hi:.0%}]")

    # direction check
    flagged = [(k, v) for k, v in rows_of(lambda k: k["direction"] in FLAG_ANSWER)
               if v[0] in OK["lenient"] and v[1] and v[1] != "can't say"]
    print("\nDirection flag vs 'Broader?' (pairs judged same/related, flag set, answer given):")
    if flagged:
        agree = sum(v[1] == FLAG_ANSWER[k["direction"]] for k, v in flagged)
        lo, hi = wilson(agree / len(flagged), len(flagged))
        print(f"  agreement {agree}/{len(flagged)} = {agree/len(flagged):.0%} [{lo:.0%} - {hi:.0%}]  (unweighted)")
        tab = Counter((k["direction"], v[1]) for k, v in flagged)
        answers = [b for b in BROADER if b != "can't say"]
        print("  " + f"{'flag':12}" + "".join(f"{b:>13}" for b in answers))
        for flag in FLAG_ANSWER:
            print("  " + f"{flag:12}" + "".join(f"{tab[(flag, b)]:>13}" for b in answers))
    else:
        print("  no flagged pairs with an answer yet")

    lost = rows_of(lambda k: k["stratum"] in below)
    lost = [(k, v) for k, v in lost if v[0] in OK["lenient"]]
    wrong = [(k, v) for k, v in rows_of(lambda k: k["stratum"] in above) if v[0] == "different"]
    for title, xs in (("Candidates judged different", wrong), ("Below-threshold pairs judged related", lost)):
        if xs:
            print(f"\n{title}:")
            for k, v in sorted(xs, key=lambda x: -x[0]["score"]):
                print(f"  #{k['#']:<4} {k['score']:.3f} {k['relation']:8} key {k['key']!r:24} "
                      f"{k['unit_a']!r} | {k['unit_b']!r}  -> {v[0]}")


# ---------------------------------------------------------------------------

def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample", help="Write the labelling workbook.")
    s.add_argument("--data", default="data/conflict/object",
                   help="folder holding the object step's output")
    s.add_argument("--pairs", default="object_pair_scores.jsonl",
                   help="object pair scores; a path, or a file name inside --data")
    s.add_argument("--norms", default="norm_keys.jsonl",
                   help="per-norm objects and units (norm_id, object, units); a path, or a file "
                        "name inside --data. Not the norm records: those go to --norm-text")
    s.add_argument("--stagea", default="data/stageA_pairs.json",
                   help="Stage A pairs; '' to skip the population check")
    s.add_argument("--allow-outside", action="store_true",
                   help="sample even if some scored pairs are not Stage A pairs")
    s.add_argument("--out", default=None)
    s.add_argument("--threshold", type=float, default=0.5)
    s.add_argument("--upper", type=float, default=0.6, help="splits head/eurovoc candidates")
    s.add_argument("--lower", type=float, default=0.4, help="splits the pairs below the threshold")
    s.add_argument("--low-split", type=float, default=None,
                   help="optional fifth level: split the lowest level here (e.g. 0.35)")
    s.add_argument("--alloc", action="append", metavar="STRATUM=N",
                   help='override one stratum size, e.g. --alloc "below | <0.40=30"')
    s.add_argument("--max-uses", type=int, default=2, help="max sampled pairs per norm")
    s.add_argument("--top-keys", type=int, default=20, help="frequent keys for the breakdown")
    s.add_argument("--known", default=None, help="CSV with columns norm_a,norm_b (diagnostic)")
    s.add_argument("--norm-text", default=None, help="norm records (JSON array or JSONL)")
    s.add_argument("--id-field", default="norm_id")
    s.add_argument("--text-field", default="text.paragraph")
    s.add_argument("--text-chars", type=int, default=1500)
    s.add_argument("--seed", type=int, default=42)
    r = sub.add_parser("review", help="Report estimates from a labelled workbook.")
    r.add_argument("--xlsx", default="data/conflict/object/pair_review.xlsx")
    args = p.parse_args()
    if args.cmd == "sample":
        args.stagea = args.stagea or None
        if args.stagea and not Path(args.stagea).is_file():
            raise SystemExit(f"Stage A file not found: {args.stagea} (pass --stagea '' to skip the check)")
        sample(args)
    else:
        review(args)


if __name__ == "__main__":
    main()